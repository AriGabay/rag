"""The citation fixtures and their manifest (U14, R33): the places the browser test expects a citation to highlight
are where the files themselves draw them.

The fixtures under ``tests/fixtures/citations/`` are synthetic (``scripts/generate_fixtures.py --only citations``),
and ``manifest.json`` records each cited place from the generator's own layout, as fractions of the shown page. Here
every box is checked against the files with pypdfium2 and pypdf only, independently of pdfplumber and of the
application's readers: a line's box holds exactly its own glyphs (pypdfium2's character boxes, put on the shown page
with the textbook transform of the CropBox and ``/Rotate``), a table cell's ruled box holds exactly its value, the
picture's box is where the image object is drawn, and every box lies on ink in the rendered page. No database.
"""

from __future__ import annotations

import json
from functools import cache
from pathlib import Path

import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
import pytest
from pypdf import PdfReader

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
CITATIONS = FIXTURES / "citations"
MANIFEST = json.loads((CITATIONS / "manifest.json").read_text(encoding="utf-8"))
DOCS = MANIFEST["documents"]
MARKER = "מסמך סינתטי לדמו"
TOLERANCE = 2.5  # points: a line box against its glyphs' loose boxes (pypdfium2's ascent is the font's, not the em)
SCALE = 2  # render scale for the ink checks


def _pdf_entries() -> list[tuple[str, dict]]:
    out = []
    for key, d in DOCS.items():
        if d["kind"] == "docx":
            continue
        out.append((key, d))
        if "replacement" in d:
            out.append((f"{key}:replacement", d["replacement"]))
    return out


PDFS = _pdf_entries()
TEXT_PDFS = [(k, d) for k, d in PDFS if d["kind"] != "pdf_scanned"]


def _places(kinds: tuple[str, ...]) -> list[tuple[str, dict, str, dict]]:
    """(doc key, entry, place name, place) for the places of the given kind: ``line`` (a drawn line), ``cells`` (a
    ruled row, cell or header), ``picture``."""
    out = []
    for key, d in TEXT_PDFS:
        for name, place in d["places"].items():
            kind = "picture" if name == "picture" else "cells" if name in ("row", "cell", "table_header") else "line"
            if kind in kinds:
                out.append((key, d, name, place))
    return out


def _path(entry: dict) -> Path:
    return FIXTURES / entry["file"]


def _shown(box, crop, rotation):
    cx0, cy0, cx1, cy1 = crop
    f = {0: lambda x, y: (x - cx0, cy1 - y), 90: lambda x, y: (y - cy0, x - cx0),
         180: lambda x, y: (cx1 - x, y - cy0), 270: lambda x, y: (cy1 - y, cx1 - x)}[rotation]
    pts = [f(x, y) for x in (box[0], box[2]) for y in (box[1], box[3])]
    return [min(p[0] for p in pts), min(p[1] for p in pts), max(p[0] for p in pts), max(p[1] for p in pts)]


@cache
def _frame(path: Path, page_no: int) -> tuple[tuple[float, float, float, float], int]:
    page = PdfReader(path).pages[page_no - 1]
    mb, cb = [float(v) for v in page.mediabox], [float(v) for v in page.cropbox]
    return (max(mb[0], cb[0]), max(mb[1], cb[1]), min(mb[2], cb[2]), min(mb[3], cb[3])), page.rotation % 360


@cache
def shown_chars(path: Path, page_no: int) -> list[tuple[str, list[float]]]:
    """Every visible character of the page with its loose box on the shown page (points, origin top-left)."""
    crop, rotation = _frame(path, page_no)
    text = pdfium.PdfDocument(str(path))[page_no - 1].get_textpage()
    out = []
    for i in range(text.count_chars()):
        ch = text.get_text_range(i, 1)
        if ch.strip():
            out.append((ch, _shown(text.get_charbox(i, loose=True), crop, rotation)))
    return out


def chars_in(path: Path, page_no: int, box: list[float], pad: float) -> list[tuple[str, list[float]]]:
    return [(ch, b) for ch, b in shown_chars(path, page_no)
            if box[0] - pad <= (b[0] + b[2]) / 2 <= box[2] + pad and box[1] - pad <= (b[1] + b[3]) / 2 <= box[3] + pad]


def union(boxes):
    return [min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes)]


@cache
def rendered(path: Path, page_no: int):
    return pdfium.PdfDocument(str(path))[page_no - 1].render(scale=SCALE).to_pil().convert("L")


def ink(path: Path, page_no: int, box: list[float]) -> int:
    x0, top, x1, bottom = (round(v * SCALE) for v in box)
    return sum(1 for v in rendered(path, page_no).crop((x0, top, x1, bottom)).getdata() if v < 128)


def glyphs(text: str) -> list[str]:
    return sorted(text.replace(" ", "").replace("|", ""))


# --- the manifest ------------------------------------------------------------------------------------------------

def test_every_fixture_is_synthetic_listed_and_present():
    files = {p.name for p in CITATIONS.iterdir() if p.suffix in (".pdf", ".docx")}
    listed = {Path(d["file"]).name for _, d in PDFS} | {Path(d["file"]).name for d in DOCS.values()}
    assert files == listed
    assert all("synthetic" in name for name in files)
    assert MARKER in MANIFEST["about"] and 0 < MANIFEST["tolerance"] <= 0.02
    for _, d in PDFS:
        assert MARKER in (PdfReader(_path(d)).metadata.title or ""), d["file"]
    from docx import Document

    assert MARKER in Document(str(_path(DOCS["docx"]))).core_properties.title


@pytest.mark.parametrize(("key", "entry"), PDFS, ids=[k for k, _ in PDFS])
def test_the_display_frame_is_the_rendered_page_and_boxes_are_its_fractions(key, entry):
    path = _path(entry)
    doc = pdfium.PdfDocument(str(path))
    assert len(doc) == entry["pages"]
    width, height = doc[0].get_size()
    # a scan is stored at 200 dpi: its page is whole pixels, a fraction of a point off the drawn A4 page
    assert (width, height) == pytest.approx((entry["display"]["width"], entry["display"]["height"]),
                                            abs=1.0 if entry["kind"] == "pdf_scanned" else 0.05)
    assert _frame(path, 1)[1] == entry["display"]["rotation"]
    for place in entry["places"].values():
        pts, frac = place["box_points"], place["box"]
        assert frac == pytest.approx([pts[0] / width, pts[1] / height, pts[2] / width, pts[3] / height], abs=0.002)
        assert 0 <= frac[0] < frac[2] <= 1 and 0 <= frac[1] < frac[3] <= 1


def test_the_queries_and_picked_passages_are_in_their_documents():
    for key, d in DOCS.items():
        if d["kind"] == "docx":
            words = " ".join(" ".join(r) for r in d["rows"])
        else:
            words = " ".join(p["text"] for p in d["places"].values())
            if "picture_table" in d:
                words += " ".join(" ".join(r) for r in d["picture_table"]["rows"])
        expects = list(d["expect"].values()) if "pick" not in d["expect"] else [d["expect"]]
        for e in expects:
            assert e["pick"]["contains"] in words, key
        hits = [w for w in d["query"].split() if w in words]
        assert len(hits) >= len(d["query"].split()) - 1, (key, d["query"])  # at most one word of the question form


# --- each place is where the file draws it -----------------------------------------------------------------------

LINES = _places(("line",))


@pytest.mark.parametrize(("key", "entry", "name", "place"), LINES, ids=[f"{k}:{n}" for k, _, n, _ in LINES])
def test_each_line_box_holds_its_own_glyphs_and_no_others(key, entry, name, place):
    path = _path(entry)
    found = chars_in(path, place["page"], place["box_points"], pad=1.0)
    assert sorted(ch for ch, _ in found) == glyphs(place["text"])
    assert union([b for _, b in found]) == pytest.approx(place["box_points"], abs=TOLERANCE)


CELLS = _places(("cells",))


@pytest.mark.parametrize(("key", "entry", "name", "place"), CELLS, ids=[f"{k}:{n}" for k, _, n, _ in CELLS])
def test_each_ruled_box_holds_exactly_its_cells_values(key, entry, name, place):
    path = _path(entry)
    box = place["box_points"]
    found = chars_in(path, place["page"], box, pad=0.0)
    assert sorted(ch for ch, _ in found) == glyphs(place["text"])
    inner = union([b for _, b in found])
    assert box[0] < inner[0] and inner[2] < box[2] and box[1] < inner[1] and inner[3] < box[3]


def test_the_repeated_number_is_in_both_the_sentence_and_the_cell_and_they_do_not_overlap():
    d = DOCS["repeated"]
    paragraph, cell, row = d["places"]["paragraph"], d["places"]["cell"], d["places"]["row"]
    assert d["number"] in paragraph["text"] and cell["text"] == d["number"] and d["number"] in row["text"]
    a, b = paragraph["box"], row["box"]
    assert a[3] < b[1]  # the sentence is above the table: neither occurrence's box touches the other
    assert row["box"][0] <= cell["box"][0] and cell["box"][2] <= row["box"][2]


def test_the_cross_page_table_continues_on_page_two_without_its_header():
    d = DOCS["cross_page"]
    header, row = d["places"]["table_header"], d["places"]["row"]
    assert header["page"] == 1 and row["page"] == 2 and d["pages"] == 2
    # the header's "(₪)" is drawn once, on page 1: nothing on page 2 writes the shekel sign
    assert "₪" in [ch for ch, _ in shown_chars(_path(d), 1)]
    assert "₪" not in [ch for ch, _ in shown_chars(_path(d), 2)]


def test_the_picture_box_is_where_the_image_is_drawn_and_holds_no_text():
    d = DOCS["mixed"]
    place = d["places"]["picture"]
    path = _path(d)
    page = pdfium.PdfDocument(str(path))[place["page"] - 1]
    crop, rotation = _frame(path, place["page"])
    images = [_shown(obj.get_bounds(), crop, rotation) for obj in page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE])]
    assert any(b == pytest.approx(place["box_points"], abs=0.5) for b in images), images
    assert chars_in(path, place["page"], place["box_points"], pad=0.0) == []
    assert ink(path, place["page"], place["box_points"]) > 1000  # the table is drawn inside it


# --- highlights lie on ink -----------------------------------------------------------------------------------------

INKED = _places(("line", "cells"))


@pytest.mark.parametrize(("key", "entry", "name", "place"), INKED, ids=[f"{k}:{n}" for k, _, n, _ in INKED])
def test_every_box_lies_on_ink_in_the_rendered_page(key, entry, name, place):
    assert ink(_path(entry), place["page"], place["box_points"]) > 10


def test_on_the_turned_page_with_an_offset_crop_box_the_sentence_box_is_on_its_ink_and_the_band_under_it_blank():
    d = DOCS["rotated"]
    path, place = _path(d), d["places"]["target"]
    assert d["display"]["rotation"] == 90
    crop, rotation = _frame(path, 1)
    mb = [float(v) for v in PdfReader(path).pages[0].mediabox]
    assert rotation == 90 and mb[0] != 0 and tuple(crop) != tuple(mb)  # a shifted MediaBox and an inset CropBox
    box = place["box_points"]
    assert ink(path, 1, box) > 100
    assert ink(path, 1, [box[0], box[3] + 4, box[2], box[3] + 10]) == 0


def test_the_scan_has_no_text_layer_and_its_cited_line_is_on_the_page_image():
    d = DOCS["scanned"]
    path = _path(d)
    assert shown_chars(path, 1) == []
    assert d["expect"]["precision"] == "page"  # OCR lines carry no position: the citation is page level
    box = d["places"]["target"]["box_points"]
    assert ink(path, 1, [box[0] - 3, box[1] - 3, box[2] + 3, box[3] + 3]) > 100


def test_the_replacement_moves_the_cited_sentence_down_and_changes_it():
    d = DOCS["replaced"]
    first, second = d["places"]["target"], d["replacement"]["places"]["target"]
    assert first["text"] != second["text"]
    assert second["box"][1] > first["box"][3]  # the original highlight is not on the replacement's sentence


# --- DOCX ----------------------------------------------------------------------------------------------------------

def test_the_docx_table_holds_the_manifest_rows_and_the_computation_inputs():
    from docx import Document

    d = DOCS["docx"]
    [table] = Document(str(_path(d))).tables
    cells = [[c.text for c in row.cells] for row in table.rows]
    assert cells == [d["headers"], *d["rows"]]
    comp = d["computation"]
    row = next(r for r in d["rows"] if r[0] == comp["row"])
    assert row[1:] == comp["inputs"]
    area, price = (int(v.replace(",", "")) for v in comp["inputs"])
    assert f"{area * price:,}" == comp["result"]
