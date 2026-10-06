"""Consistency checks between the committed synthetic fixtures and ground_truth.yaml (U17).

Runs on the host without Docker: it only reads the committed files.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
import yaml
from pypdf import PdfReader

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
TRUTH = yaml.safe_load((FIXTURES / "ground_truth.yaml").read_text(encoding="utf-8"))
DOCS = {d["id"]: d for d in TRUTH["documents"]}
RECORDS = [(d, r) for d in TRUTH["documents"] for r in d["records"]]
MARKER = "מסמך סינתטי לדמו"
TABLE_HEADERS = [
    "כתובת",
    "גוש/חלקה",
    "תאריך עסקה",
    "סוג נכס",
    "חדרים",
    "שטח (מ״ר)",
    "סוג שטח",
    "מחיר (₪)",
    "מחיר למ״ר (₪)",
]


@pytest.mark.parametrize("doc", TRUTH["documents"], ids=lambda d: d["id"])
def test_every_ground_truth_file_exists_and_is_marked_synthetic(doc):
    assert "synthetic" in doc["filename"]
    assert (FIXTURES / doc["filename"]).is_file()


def test_no_untracked_fixture_documents():
    listed = {d["filename"] for d in TRUTH["documents"]}
    on_disk = {p.name for p in FIXTURES.iterdir() if p.suffix in {".pdf", ".docx"}}
    assert on_disk == listed


@pytest.mark.parametrize(
    "doc",
    [d for d in TRUTH["documents"] if d["kind"] in {"pdf_digital", "pdf_visual", "pdf_scanned"}],
    ids=lambda d: d["id"],
)
def test_pdf_opens_and_page_count_matches(doc):
    reader = PdfReader(FIXTURES / doc["filename"])
    assert not reader.is_encrypted
    assert len(reader.pages) == doc["page_count"]
    for section in doc["sections"]:
        assert 1 <= section["page"] <= doc["page_count"]
    if doc["kind"] != "pdf_scanned":
        text = "".join(page.extract_text() for page in reader.pages)
        assert text.strip(), "digital PDFs must have a text layer"


def test_scanned_pdf_has_no_text_layer():
    reader = PdfReader(FIXTURES / DOCS["D2"]["filename"])
    assert all(not (page.extract_text() or "").strip() for page in reader.pages)


def test_encrypted_pdf_reports_encrypted_and_opens_with_password():
    doc = DOCS["BAD_encrypted"]
    reader = PdfReader(FIXTURES / doc["filename"])
    assert reader.is_encrypted
    assert reader.decrypt(doc["password"])
    assert len(reader.pages) == doc["page_count"]
    assert doc["records"] == []


def test_truncated_pdf_fails_to_open():
    doc = DOCS["BAD_truncated"]
    with pytest.raises(Exception):  # noqa: B017 - any parse error is the expected outcome
        reader = PdfReader(FIXTURES / doc["filename"])
        _ = [page.extract_text() for page in reader.pages]
    assert doc["records"] == []


def test_docx_contains_marker_and_table_header():
    import docx

    document = docx.Document(FIXTURES / DOCS["D11"]["filename"])
    assert any(MARKER in p.text for p in document.paragraphs)
    assert [c.text for c in document.tables[0].rows[0].cells] == TABLE_HEADERS
    assert len(document.tables[0].rows) - 1 == len(DOCS["D11"]["tables"][0]["rows"])


def test_record_ids_unique():
    ids = [r["id"] for _, r in RECORDS]
    assert len(ids) == len(set(ids))


def test_office_a_has_enough_transaction_records():
    count = sum(
        1
        for d, r in RECORDS
        if d["office"] == "A" and d["version_of"] is None and r["data_kind"] == "transaction_price"
    )
    assert count >= 32


def test_every_report_has_one_appraised_value_record():
    for doc in TRUTH["documents"]:
        if doc["kind"] in {"encrypted", "truncated"}:
            continue
        kinds = [r["data_kind"] for r in doc["records"]]
        assert kinds.count("appraised_value") == 1, doc["id"]


def test_stated_price_per_sqm_matches_price_over_area_except_d6_conflict():
    conflicts = []
    for _doc, r in RECORDS:
        if r["price_per_sqm_stated"] is None or r["area"] is None:
            continue
        computed = Decimal(r["price"]) / Decimal(r["area"])
        stated = Decimal(r["price_per_sqm_stated"])
        if abs(stated - computed) / computed > Decimal("0.005"):
            conflicts.append(r["id"])
            assert r["price_per_sqm_conflict"] is True
        else:
            assert r["price_per_sqm_conflict"] is False
    assert conflicts == ["D6-T00"]


def test_d9_has_thirty_rows_and_records():
    table = DOCS["D9"]["tables"][0]
    assert len(table["rows"]) == 30
    assert sum(r["data_kind"] == "transaction_price" for r in DOCS["D9"]["records"]) == 30


def test_d3_table_crosses_pages_without_repeated_header():
    table = DOCS["D3"]["tables"][0]
    assert len(table["rows"]) >= 14
    assert table["page_start"] < table["page_end"]
    assert table["header_repeated_on_continuation"] is False
    assert {row["page"] for row in table["rows"]} == {table["page_start"], table["page_end"]}


def test_tables_and_records_are_consistent():
    for doc in TRUTH["documents"]:
        for table in doc["tables"]:
            assert table["headers"] == TABLE_HEADERS
            for row in table["rows"]:
                assert len(row["cells"]) == len(TABLE_HEADERS)
        for r in doc["records"]:
            assert r["currency"] == "ILS"
            assert r["area_type"] in {None, "net", "gross", "registered", "equivalent"}
            if r["data_kind"] == "transaction_price":
                row = doc["tables"][r["table_index"]]["rows"][r["row_index"]]
                assert r["page"] == row["page"]


def test_dedup_expectations():
    certain = [sorted(o["record"] for o in g["occurrences"]) for g in TRUTH["dedup"]["certain"]]
    uncertain = [sorted(o["record"] for o in g["occurrences"]) for g in TRUTH["dedup"]["uncertain"]]
    assert certain == [["D1-T01", "D4-T00"]]
    assert uncertain == [["D1-T00", "D4-T01"]]
    by_id = {r["id"]: r for _, r in RECORDS}
    a, b = by_id["D1-T00"], by_id["D4-T01"]
    assert (a["area"], a["area_type"]) == ("95", "net")
    assert (b["area"], b["area_type"]) == ("110", "gross")


def test_d1v2_changes_exactly_one_price():
    old = {r["row_index"]: r for r in DOCS["D1"]["records"] if r["row_index"] is not None}
    new = {r["row_index"]: r for r in DOCS["D1v2"]["records"] if r["row_index"] is not None}
    changed = [i for i in old if old[i]["price"] != new[i]["price"]]
    assert changed == [3]


def test_content_facts_point_at_known_documents():
    for fact in TRUTH["content_facts"]:
        doc = DOCS[fact["document"]]
        if doc["page_count"] is None:
            assert fact["page"] is None
        else:
            assert 1 <= fact["page"] <= doc["page_count"]
