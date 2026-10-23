"""Round 21 acceptance with the real model (U10: R1-R3, R5, R6, R11, R12; AE1, AE3, AE4 (the several-lots shape), AE6, AE8): the shapes the
round-21 binding fixes handle, shown on the configured model and judge (``gpt-6-luna`` with the repository's settings, no override; the
key comes from the repository's ``.env`` through the round-10 ``real_key`` fixture and is never printed). Synthetic documents only (every
name, address and figure invented). No binary fixture is added: the documents of a1-a3 are DOCX files written into the test's own
temporary directory (never into the repository), a4 and a5 use the round-20 PDFs of ``tests/fixtures/round20``
(``scripts/generate_fixtures.py --only round20``; the contents are the manifest's), ingested through the worker into a synthetic office
of the ``rag_test`` database.

Every check reads the BINDING RECORD where it can, not only the text: the registry of the spans the server wrote (``server_spans``: kind,
offsets, ``claim_ids`` and ``datum`` - round 21 U2, KTD1), the item ledger the datum ids point into, the verification rows of the
diagnostics (the claims removed, cut or withheld, with their claim ids), the captured drafts (round 21 U1) and the sources as served. The
text is read only for what the record cannot say (the place a value stands, the period a clause names).

GATED (all three repetitions must pass; a repetition in which the model writes no such shape is recorded as ``shape not produced`` and
counted separately - it is not a pass):

- ``a1`` two data with their own conditions in one sentence (AE1): both data reach the answer; every condition the server wrote inside
  the answer names its datum (``datum`` or ``claim_ids``), the datum is the one whose condition it is (the ledger item the span names and
  the nearest value before the span agree), and no condition stands after the other datum; no condition is hidden;
- ``a2`` a capitalised total next to a monthly rate in one table (AE3): both reach the answer; no server qualifier that names a period
  is attached to the capitalised total (by its datum, its position, and the claim rows), and the monthly rate keeps its period;
- ``a3`` several lots with an identical bullet under each lot heading (the tidy-up of U4, AE4): the bullet of each lot is kept (as many
  as the last verified draft held), each lot's own value stands under its own lot and never under the other, no removal calls a lot's own
  datum another property's and no server sentence says "another property" about lots of one property;
- ``a4`` a follow-up question in the same conversation (AE8): the first answer states the group's payment condition (the shape); the
  second answer carries no condition note of the first - no server span, no item condition span of the new ledger, no claim id shared
  with the first answer's spans, and no stated line that puts the condition beside the new group's datum;
- ``a5`` a reversed-stored left-to-right table (AE6, built as ``test_round21_sequences`` builds it): the answer gives the value under the
  right column label, the sources it cites serve the row in drawn order with each cell paired with its label, no source serves it
  reversed, and the value is not removed by verification.

Each sequence is asked in a **new conversation, exactly three times**; every repetition is recorded before anything is asserted - a
failure in any repetition is reported and a later pass never hides it. The checks are by behaviour and never by wording (a pattern part
is matched, so a Hebrew definite article or prefix does not change the outcome).

Recorded per repetition (``real_documents/eval/round21/real-model/``, gitignored: the documents are synthetic, so the answers are kept
for diagnosis): the whole turn (final text, rounds, tools, item ledger, the spans the server wrote with kind, text, claim ids and datum,
the removals and the verification rows, the captured drafts, usage), the pass or fail of each check, and the stage that produced the part.
A summary file per sequence is written when the sequence ends - before anything is asserted - and one for the run at the end of the module.

Opt in with ``-m real_model``; skipped without a key. Host only, against ``rag_test`` (``tests/conftest.py`` guard)."""

from __future__ import annotations

import io
import json
import os
import re
import time
from datetime import UTC, datetime
from pathlib import Path

import pytest
from sqlalchemy import text

from app.config import get_settings
from app.db import tenant_tx
from tests.conftest import login
from tests.factories import make_office
from tests.integration.test_round10_real_model import new_conversation, upload
from tests.integration.test_round10_real_model import real_key as real_key  # noqa: F401 - the key fixture
from tests.integration.test_round17_real_model import answered, cites_document, has
from tests.integration.test_round18_real_model import figure
from tests.integration.test_round19_real_model import REMOVAL_SENTENCE, Turn19, base_info, cells_of, stated
from tests.integration.test_round21_sequences import (  # the AE6 builders, reused as they are
    QUARTER_PAIRS,
    QUARTERS,
    REVENUE_QUESTION,
    pairs_in_drawn_order,
    reverse_stored_table,
)

pytestmark = [pytest.mark.real_model, pytest.mark.db]

ROOT = Path(__file__).resolve().parents[3]
FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
MANIFEST = json.loads((FIXTURES / "round20" / "manifest.json").read_text(encoding="utf-8"))
DOCS = MANIFEST["documents"]
SYNTHETIC_MARKER = MANIFEST["synthetic_marker"]
LOTS = DOCS["lot_groups"]
RESULTS_DIR = ROOT / "real_documents" / "eval" / "round21" / "real-model"
STAMP = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
RESULTS_FILE = RESULTS_DIR / f"results-{STAMP}.json"
SUMMARY_FILE = RESULTS_DIR / f"summary-{STAMP}.json"
RESULTS: list[dict] = []
ATTEMPTS = 3  # exactly three, each in a new conversation
STARTED = time.perf_counter()
MODEL = "gpt-6-luna (settings)"

ADMIN = "admin-r21@example.test"

# the spans the server writes INSIDE the answer's sentences and cells (the ones that must name a datum, KTD1) and those it writes in the gap
# area after the answer (the lines about items not given, the one uncertainty line, the calculation's caveat)
INLINE_KINDS = ("qualifier", "scope_note", "partial_mark", "citation", "marker")
GAP_KINDS = ("item_gap", "gap_paragraph", "uncertainty_line", "conditional_note")
GONE_ACTIONS = ("removed", "withheld", "cut")  # the verification rows that take a claim (or a part of it) out of the text
CLAIM_ID = re.compile(r"k[0-9a-f]{12}")
BULLET = re.compile(r"^\s*(?:[-*•]|\d+[.)])\s+")
MARKUP = re.compile(r"[*`‎‏‪-‮⁦-⁩]")
SENTENCE_END = re.compile(r"(?<=[.!?])\s+")
ANOTHER_PROPERTY_WORDS = re.compile(r"נכס\s+אחר")
ANOTHER_PROPERTY = "wrong_subject"


# --- the synthetic documents (written into the test's own temporary directory, never into the repository) ------------------------

def docx_21(title: str, blocks: list[tuple[str, object]]) -> bytes:
    """A DOCX of ``blocks`` (``("h1", text)``, ``("p", text)`` or ``("table", [[cell, ...], ...])``) under ``title``, with Word heading
    styles so the reader finds the structure, and the manifest's synthetic mark as its first paragraph."""
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"{SYNTHETIC_MARKER} — כל הנתונים בדויים")
    for kind, body in blocks:
        if kind == "h1":
            doc.add_heading(str(body), 1)
        elif kind == "table":
            rows = list(body)  # type: ignore[arg-type]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, value in enumerate(cells):
                    table.cell(r, c).text = value
        else:
            doc.add_paragraph(str(body))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


PURPOSE = [("h1", "1. מטרת השומה"), ("p", "שומה זו נערכה לצורך הדגמה בלבד (מסמך סינתטי).")]
# a1: two areas, each with a condition of its own in the note column (the unit test's TWO_CONDITIONS_ROWS)
CLAUSE_A, CLAUSE_B = "הזכויות חלות רק על פי תכנית מאושרת", "השטח חל רק לאחר פינוי המגרש"
TWO_CONDITIONS_DOC = ("שומת הדגמה לנכס ברחוב הגפן 7 — זכויות בנייה",
                      [*PURPOSE, ("h1", "3. זכויות הבנייה"), ("p", "הטבלה שלהלן מפרטת את השטחים בנכס ואת ההערה לכל אחד מהם."),
                       ("table", [["סוג שטח", "שטח (מ״ר)", "הערה"], ["שטח עיקרי", "1,200", CLAUSE_A], ["שטח שירות", "300", CLAUSE_B]])])
# a2: a monthly column with a monthly rent row and a capitalised total row (the unit test's INCOME_TABLE)
INCOME_DOC = ("שומת הדגמה למרכז מסחרי ברחוב השקמה 3 — תמהיל הכנסות",
              [*PURPOSE, ("h1", "4. תמהיל הכנסות"), ("p", "הטבלה שלהלן מציגה את מרכיבי ההכנסה ואת השווי המהוון של הנכס."),
               ("table", [["רכיב", "סכום חודשי (₪)"], ["שכר דירה בסיסי", "55,000"], ["שווי מהוון", "6,600,000"]])])
# a3: two lots with the SAME area and different land values (the unit test's SAME_BULLET_SOURCE with a value each, to tell the lots apart)
LOTS_DOC = ("שומת הדגמה למגרשים ברחוב הגפן 7 — שטח ושווי הקרקע",
            [*PURPOSE, ("h1", "3. המגרשים"), ("h1", "מגרש 101"), ("p", "שטח המגרש 800 מ״ר. שווי הקרקע במגרש 3,100,000 ₪."),
             ("h1", "מגרש 102"), ("p", "שטח המגרש 800 מ״ר. שווי הקרקע במגרש 4,200,000 ₪.")])


# --- what each sequence looks for -----------------------------------------------------------------------------------------

# a1
MAIN, SERVICE = figure("1200"), figure("300")
A1_VALUES = {"main": MAIN, "service": SERVICE}
A1_CONDITIONS = {"main": re.compile(r"מאושרת"), "service": re.compile(r"פינוי")}
A1_LABELS = {"main": re.compile(r"עיקרי"), "service": re.compile(r"שירות")}
# a2
RENT, TOTAL = figure("55000"), figure("6600000")
A2_VALUES = {"rent": RENT, "total": TOTAL}
A2_LABELS = {"rent": re.compile(r"שכר\s+דירה|שכירות"), "total": re.compile(r"מהוון|היוון")}
PERIOD = re.compile(r"חודש|monthly", re.IGNORECASE)  # "monthly", whatever the prefix or the form (חודשי / לחודש / בחודש)
# a3
LOT_AREA = figure("800")
LOT_VALUE = {"101": figure("3100000"), "102": figure("4200000")}
LOT_MENTION = {"101": re.compile(r"(?<![\d,.])101(?![\d,])"), "102": re.compile(r"(?<![\d,.])102(?![\d,])")}
# a4 / R20b: the two groups by their areas, the footnote's condition (the land authority's approval)
(GROUP_A_ROW, GROUP_B_ROW) = LOTS["rows"]
AREA_A, AREA_B = figure(GROUP_A_ROW[2].replace(",", "")), figure(GROUP_B_ROW[2].replace(",", ""))
CONDITION = re.compile(r"מינהל|מנהל\s+מקרקעי|רמ״י|מקרקעי\s+ישראל")
SERVER_CONDITION_NOTE = re.compile(r"התנאי\s+לפי\s+המקור")
# a5: the quarters of the English table of R20f (Revenue: Q1 2024 100, Q2 2024 200, Annual 300)
QUARTER_VALUES = {"Q1": figure("100"), "Q2": figure("200"), "Annual": figure("300")}
QUARTER_LABELS = {"Q1": re.compile(r"Q1|רבעון\s+(?:ה)?ראשון|ראשון"), "Q2": re.compile(r"Q2|רבעון\s+(?:ה)?שני|שני"),
                  "Annual": re.compile(r"Annual|שנתי|שנה|סה״כ|סך\s+הכל")}


# --- environment ----------------------------------------------------------------------------------------------------------

@pytest.fixture
def office(db, client, monkeypatch, real_key):  # noqa: F811 - the key fixture is the imported one
    monkeypatch.setattr(get_settings(), "chat_run_inline", True)
    o = make_office(db, "משרד בדיקה סינתטי", ADMIN)
    with tenant_tx(o.ctx()) as conn:
        conn.execute(text("UPDATE office_settings SET cloud_llm_enabled = true"))
    login(client, ADMIN)
    o.client, o.ids = client, {}
    return o


def ingest_pdf(office, monkeypatch, spec: dict) -> str:
    """A round-20 PDF fixture read by the worker; the document id."""
    key = spec["file"]
    office.ids[key] = upload(office, FIXTURES / spec["file"], "appraisal", monkeypatch)
    return office.ids[key]


def ingest_docx(office, monkeypatch, tmp_path: Path, name: str, doc: tuple[str, list]) -> str:
    """A synthetic DOCX written in the test's temporary directory and read by the worker; the document id."""
    path = tmp_path / name
    path.write_bytes(docx_21(*doc))
    office.ids[name] = upload(office, path, "appraisal", monkeypatch)
    return office.ids[name]


# --- turns and records ----------------------------------------------------------------------------------------------------

def flush() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    RESULTS_FILE.write_text(json.dumps({"model": MODEL, "attempts": RESULTS}, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


class Turn21(Turn19):
    """The round-19 turn record with the BINDING RECORD of the answer: every server span with its ``claim_ids`` and ``datum`` (round 21 U2,
    KTD1), the gap lines cut with the uncertainty line (round 21 U2, KTD2) included, the verification rows of the diagnostics (a claim
    removed, cut, withheld, qualified or cited, with its claim id) and the captured drafts (round 21 U1)."""

    def __init__(self, c, m: dict, question: str, wall: float) -> None:
        super().__init__(c, m, question, wall)
        d = c.get(f"/api/chat/messages/{m['id']}/diagnostics")
        body = d.json() if d.status_code == 200 else {}
        self.claim_rows = [{k: r.get(k) for k in ("claim_id", "action", "check", "failure_kind", "verdict", "text", "remaining", "reason")}
                           for r in body.get("removed") or []]
        self.drafts = [{"markdown": x.get("markdown") or ""} for x in body.get("drafts") or []]

    @property
    def spans(self) -> list[dict]:
        md = self.markdown
        return [{"kind": s.get("kind"), "start": s["start"], "end": s["end"], "text": md[s["start"]:s["end"]],
                 "claim_ids": [c for c in s.get("claim_ids") or [] if isinstance(c, str)], "datum": s.get("datum") or ""}
                for s in self.answer.get("server_spans") or [] if isinstance(s.get("start"), int) and isinstance(s.get("end"), int)]

    @property
    def inline_spans(self) -> list[dict]:
        return [s for s in self.spans if s["kind"] in INLINE_KINDS]

    @property
    def gap_lines(self) -> list[str]:
        out = []
        for s in self.spans:
            if s["kind"] in GAP_KINDS:
                out += [re.sub(r"\s+", " ", ln).strip() for ln in s["text"].splitlines() if ln.strip()]
        return out

    @property
    def stated_lines(self) -> list[str]:
        """The answer's lines without the spans of the gap kinds (the uncertainty line included): what the answer states by itself."""
        md, cuts = self.markdown, sorted((s["start"], s["end"]) for s in self.spans if s["kind"] in GAP_KINDS)
        out, pos = [], 0
        for a, b in cuts:
            out.append(md[pos:a])
            pos = max(pos, b)
        out.append(md[pos:])
        return [re.sub(r"\s+", " ", MARKUP.sub("", ln)).strip() for ln in "".join(out).splitlines() if ln.strip()]

    @property
    def body_lines(self) -> list[tuple[int, str]]:
        """The raw lines of the answer (offset, text) outside the gap-kind spans, markup kept (a bullet is told by it)."""
        cuts = sorted((s["start"], s["end"]) for s in self.spans if s["kind"] in GAP_KINDS)
        return [(m.start(), m.group()) for m in re.finditer(r"[^\n]+", self.markdown)
                if m.group().strip() and not any(a <= m.start() < b for a, b in cuts)]

    def item(self, item_id: str) -> dict | None:
        return next((i for i in self.items if i["id"] == item_id), None)

    def record(self) -> dict:
        return super().record() | {"claim_rows": self.claim_rows, "drafts": self.drafts}


def ask(c, cid: str, question: str, rec: dict) -> Turn21:
    """One turn through the API; recorded (and flushed) before anything is asserted about it."""
    started = time.perf_counter()
    r = c.post(f"/api/chat/conversations/{cid}/messages", json={"content": question})
    assert r.status_code == 200, r.text
    m = c.get(f"/api/chat/messages/{r.json()['assistant']['id']}").json()
    t = Turn21(c, m, question, time.perf_counter() - started)
    rec["turns"].append(t.record())
    flush()
    assert m["status"] == "done", (m.get("status"), m.get("error"))
    return t


def check(rec: dict, name: str, fn) -> None:
    """Run one named assertion of a repetition and record its outcome; a failure never stops the repetition's other checks."""
    try:
        fn()
    except AssertionError as exc:
        rec["checks"][name] = {"outcome": "fail", "failure": str(exc)[:2000]}
    else:
        rec["checks"][name] = {"outcome": "pass"}
    flush()


def sequence_summary(sequence: str) -> dict:
    runs = [r for r in RESULTS if r["sequence"] == sequence]
    attempts = []
    for r in runs:
        info = r.get("info") or {}
        turns = r.get("turns") or [{}]
        attempts.append({"attempt": r["attempt"], "outcome": r["outcome"],
                         "checks": {k: v["outcome"] for k, v in r["checks"].items()},
                         "failures": {k: v["failure"] for k, v in r["checks"].items() if v["outcome"] == "fail"},
                         "stage": info.get("stage"), "info": {k: v for k, v in info.items() if k != "stage"},
                         "final_text": turns[-1].get("markdown"), "server_spans": turns[-1].get("server_spans"),
                         "rounds": turns[-1].get("rounds"), "tools": turns[-1].get("tools"),
                         "turns": [{"question": t.get("question"), "markdown": t.get("markdown"), "server_spans": t.get("server_spans"),
                                    "claim_rows": t.get("claim_rows")} for t in turns]})
    passed = sum(a["outcome"] == "pass" for a in attempts)
    not_produced = sum(1 for a in attempts if a["info"].get("shape_produced") is False)
    return {"model": MODEL, "sequence": sequence, "attempts": attempts, "passed_runs": f"{passed}/{len(attempts)}",
            "shape_not_produced_runs": not_produced, "passed": len(attempts) == ATTEMPTS and passed == ATTEMPTS}


def run_sequence(sequence: str, checks: tuple[str, ...], office, body) -> None:
    """``body(office, client, conversation_id, record)`` in a new conversation, exactly ``ATTEMPTS`` times; every repetition is recorded and
    none is skipped after a failure. The summary file of the sequence is written before anything is asserted. The test fails when any
    check of any repetition failed (or never ran because a turn did not finish), with every failure in the message."""
    failures: list[str] = []
    for n in range(1, ATTEMPTS + 1):
        rec = {"sequence": sequence, "attempt": n, "turns": [], "checks": {}, "info": {}, "outcome": None}
        RESULTS.append(rec)
        try:
            body(office, office.client, new_conversation(office.client), rec)
        except AssertionError as exc:
            rec["aborted"] = str(exc)[:2000]
        for name in checks:
            rec["checks"].setdefault(name, {"outcome": "fail", "failure": "never ran: " + rec.get("aborted", "?")})
        rec["outcome"] = "pass" if all(v["outcome"] == "pass" for v in rec["checks"].values()) else "fail"
        flush()
        failures += [f"attempt {n} / {name}: {v['failure']}" for name, v in rec["checks"].items() if v["outcome"] == "fail"]
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    (RESULTS_DIR / f"summary-{sequence}-{STAMP}.json").write_text(
        json.dumps(sequence_summary(sequence), ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    assert not failures, "\n".join(failures)


# --- shared readers -------------------------------------------------------------------------------------------------------

def nearest_before(md: str, pos: int, anchors: dict[str, re.Pattern]) -> str | None:
    """The key of the anchor whose last match ends nearest before ``pos`` in ``md`` (the datum a server addition written at ``pos``
    stands right after), None when no anchor stands before it."""
    best, key = -1, None
    for k, pat in anchors.items():
        for m in pat.finditer(md, 0, pos):
            if m.end() > best:
                best, key = m.end(), k
    return key


def item_is(item: dict | None, key: str, values: dict[str, re.Pattern], labels: dict[str, re.Pattern]) -> bool:
    """Whether a ledger item is the datum ``key`` - by its value or by its label."""
    return bool(item) and (has(item.get("value") or "", values[key]) or has(item.get("label") or "", labels[key]))


def clause_of(line: str, anchor: re.Pattern, others: list[re.Pattern]) -> str:
    """The clause of a sentence that holds ``anchor``: from the end of the previous other datum before it to the start of the next other
    datum after it (the whole line when it stands alone). A table row is read whole (its header is added by the caller)."""
    m = anchor.search(line)
    if m is None:
        return ""
    if line.lstrip().startswith("|"):
        return line
    found = [o for o in (p.search(line) for p in others) if o]
    before = [o.end() for o in found if o.end() <= m.start()]
    after = [o.start() for o in found if o.start() >= m.end()]
    return line[max(before, default=0):min(after, default=len(line))]


def clauses(t: Turn21, anchor: re.Pattern, others: list[re.Pattern]) -> list[str]:
    """For each stated line with ``anchor``: its clause, with the header line of its table when it is a table row."""
    lines, out = t.stated_lines, []
    for i, ln in enumerate(lines):
        if not has(ln, anchor):
            continue
        clause = clause_of(ln, anchor, others)
        if ln.lstrip().startswith("|"):
            start = i
            while start > 0 and lines[start - 1].lstrip().startswith("|"):
                start -= 1
            clause += " " + (lines[start] if start != i else "")
        out.append(clause)
    return out


def span_view(t: Turn21, spans: list[dict], values: dict[str, re.Pattern], labels: dict[str, re.Pattern]) -> list[dict]:
    """The record of the given spans: kind, text, claim ids, datum, the ledger item the datum names (its value and label) and the datum
    whose value stands nearest before the span in the text."""
    return [{"kind": s["kind"], "text": s["text"][:200], "claim_ids": s["claim_ids"], "datum": s["datum"],
             "datum_item": (lambda i: {"id": i["id"], "label": i["label"], "value": i["value"]} if i else None)(t.item(s["datum"])) if s["datum"] else None,
             "datum_is": [k for k in values if item_is(t.item(s["datum"]) if s["datum"] else None, k, values, labels)],
             "nearest_value_before": nearest_before(t.markdown, s["start"], values)} for s in spans]


def verification_rows_with(t: Turn21, anchor: re.Pattern) -> list[dict]:
    """The verification rows that took a claim (or a part of it) out of the text and whose original text holds ``anchor``."""
    return [r for r in t.claim_rows if r["action"] in GONE_ACTIONS and has(r["text"] or "", anchor)]


def shape_not_produced(t: Turn21, ok: bool, what: str) -> None:
    assert ok, ("SHAPE NOT PRODUCED: " + what, t.markdown[:500])


def _all(reached: dict[str, bool], t: Turn21) -> None:
    missing = [n for n, ok in reached.items() if not ok]
    assert not missing, ("a part did not reach the answer", missing, t.markdown[:500])


# --- a1: two data with their own conditions in one sentence (AE1) ----------------------------------------------------------------

A1_QUESTION = "מהם השטח העיקרי ושטח השירות בנכס ברחוב הגפן 7, ומהו התנאי של כל אחד מהם? ענה במשפט אחד."


def test_a1_two_data_with_their_own_conditions_keep_each_condition_in_its_own_clause_three_times(office, monkeypatch, tmp_path):
    doc = ingest_docx(office, monkeypatch, tmp_path, "R21_synthetic_two_conditions.docx", TWO_CONDITIONS_DOC)

    def one_sentence(t: Turn21) -> bool:
        return any(has(s, MAIN) and has(s, SERVICE) for ln in t.stated_lines if not ln.startswith("|") for s in SENTENCE_END.split(ln))

    def condition_spans(t: Turn21) -> list[tuple[dict, str]]:
        """Each inline server span that carries one of the two conditions, with the datum whose condition it is."""
        out = []
        for s in t.inline_spans:
            owners = [k for k, pat in A1_CONDITIONS.items() if pat.search(s["text"])]
            if owners:
                out.append((s, owners[0] if len(owners) == 1 else "both"))
        return out

    def each_condition_bound_to_its_own_datum(t: Turn21) -> None:
        bad = []
        for s, owner in condition_spans(t):
            other = "service" if owner == "main" else "main"
            if owner == "both":
                bad.append(("one addition carries both conditions", s["text"][:160]))
                continue
            if not (s["datum"] or s["claim_ids"]):
                bad.append(("an addition inside the answer names no datum", s["text"][:160]))
            if s["claim_ids"] and not all(CLAIM_ID.fullmatch(c) for c in s["claim_ids"]):
                bad.append(("a claim id is not a claim id", s["claim_ids"]))
            item = t.item(s["datum"]) if s["datum"] else None
            if item and item_is(item, other, A1_VALUES, A1_LABELS) and not item_is(item, owner, A1_VALUES, A1_LABELS):
                bad.append(("the span names the other datum's ledger item", s["datum"], item["label"], item["value"], s["text"][:160]))
            nearest = nearest_before(t.markdown, s["start"], A1_VALUES)
            if nearest not in (owner, None):
                bad.append(("the condition stands after the other datum's value", owner, s["text"][:160]))
        assert not bad, ("a condition the server wrote is not bound to its own datum", bad[:4])

    def conditions_not_hidden(t: Turn21) -> None:
        missing = [k for k, pat in A1_CONDITIONS.items() if not has(t.markdown, pat)]
        assert not missing, ("a condition of the source does not stand anywhere in the answer", missing, t.markdown[:500])

    def body(o, c, cid, rec):
        t = ask(c, cid, A1_QUESTION, rec)
        reached = {"main": stated(t, MAIN), "service": stated(t, SERVICE)}
        rec["info"] = base_info(t, reached, {"main": t.part_stage(MAIN), "service": t.part_stage(SERVICE)})
        text_places = []  # where the conditions stand in the whole text (the model's own words included): reported, not asserted
        for k, pat in A1_CONDITIONS.items():
            for m in pat.finditer(t.markdown):
                text_places.append((k, nearest_before(t.markdown, m.start(), A1_VALUES), any(s["start"] <= m.start() < s["end"] for s in t.spans)))
        rec["info"] |= {"shape_produced": one_sentence(t), "condition_spans": span_view(t, [s for s, _ in condition_spans(t)], A1_VALUES, A1_LABELS),
                        "inline_spans": span_view(t, t.inline_spans, A1_VALUES, A1_LABELS),
                        "uncertainty_lines": [s["text"][:240] for s in t.spans if s["kind"] == "uncertainty_line"],
                        "condition_text_places": text_places, "removed_rows": [(r["action"], r["failure_kind"], (r["text"] or "")[:120]) for r in t.claim_rows
                                                                                if r["action"] in GONE_ACTIONS]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "shape_produced", lambda: shape_not_produced(t, one_sentence(t), "the two data are not stated in one sentence"))
        check(rec, "both_data_present", lambda: _all(reached, t))
        check(rec, "each_condition_bound_to_its_own_datum", lambda: each_condition_bound_to_its_own_datum(t))
        check(rec, "conditions_not_hidden", lambda: conditions_not_hidden(t))
        check(rec, "no_datum_removed", lambda: _no(verification_rows_with(t, MAIN) + verification_rows_with(t, SERVICE), "a datum was removed by verification"))

    run_sequence("a1_two_data_own_conditions",
                 ("answered", "cited", "shape_produced", "both_data_present", "each_condition_bound_to_its_own_datum", "conditions_not_hidden",
                  "no_datum_removed"), office, body)


def _no(found: list, message: str) -> None:
    assert not found, (message, found[:3])


# --- a2: a capitalised total next to a monthly rate (AE3) ------------------------------------------------------------------------

A2_QUESTION = "מהו שכר הדירה הבסיסי ומהו השווי המהוון של הנכס ברחוב השקמה 3?"


def test_a2_a_capitalised_total_carries_no_monthly_qualifier_and_the_monthly_rate_keeps_its_period_three_times(office, monkeypatch, tmp_path):
    doc = ingest_docx(office, monkeypatch, tmp_path, "R21_synthetic_income_mix.docx", INCOME_DOC)

    def total_has_no_monthly_server_qualifier(t: Turn21) -> None:
        bad = []
        for s in t.inline_spans:
            if not PERIOD.search(s["text"]):
                continue
            item = t.item(s["datum"]) if s["datum"] else None
            by_datum = item_is(item, "total", A2_VALUES, A2_LABELS) and not item_is(item, "rent", A2_VALUES, A2_LABELS)
            by_place = nearest_before(t.markdown, s["start"], A2_VALUES) == "total"
            if by_datum or by_place:
                bad.append(("a server addition naming a period is attached to the capitalised total", by_datum, by_place, s["datum"],
                            s["claim_ids"], s["text"][:160]))
        assert not bad, ("the capitalised total carries a monthly qualifier the server wrote", bad[:3])

    def monthly_rate_keeps_its_period(t: Turn21) -> None:
        found = clauses(t, RENT, [TOTAL])
        assert found, ("the monthly rate is not stated", t.markdown[:400])
        assert any(PERIOD.search(c) for c in found), ("the monthly rate stands without its period", found[:3])

    def body(o, c, cid, rec):
        t = ask(c, cid, A2_QUESTION, rec)
        reached = {"rent": stated(t, RENT), "total": stated(t, TOTAL)}
        rec["info"] = base_info(t, reached, {"rent": t.part_stage(RENT), "total": t.part_stage(TOTAL)})
        rec["info"] |= {"inline_spans": span_view(t, t.inline_spans, A2_VALUES, A2_LABELS),
                        "period_spans": span_view(t, [s for s in t.inline_spans if PERIOD.search(s["text"])], A2_VALUES, A2_LABELS),
                        "rent_clauses": clauses(t, RENT, [TOTAL]), "total_clauses_model_text": clauses(t, TOTAL, [RENT]),
                        "claim_rows_of_total": [(r["claim_id"], r["action"], r["verdict"], (r["text"] or "")[:120]) for r in t.claim_rows
                                                if has(r["text"] or "", TOTAL)]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "both_values_present", lambda: _all(reached, t))
        check(rec, "total_has_no_monthly_server_qualifier", lambda: total_has_no_monthly_server_qualifier(t))
        check(rec, "monthly_rate_keeps_its_period", lambda: monthly_rate_keeps_its_period(t))
        check(rec, "no_datum_removed", lambda: _no(verification_rows_with(t, RENT) + verification_rows_with(t, TOTAL),
                                                    "a datum was removed by verification"))

    run_sequence("a2_capitalised_total_and_monthly_rate",
                 ("answered", "cited", "both_values_present", "total_has_no_monthly_server_qualifier", "monthly_rate_keeps_its_period",
                  "no_datum_removed"), office, body)


# --- a3: several lots with an identical bullet under each lot heading (the tidy-up of U4, AE4) -------------------------------------

A3_QUESTION = ("מהם השטח ושווי הקרקע של כל אחד ממגרשים 101 ו-102 בנכס ברחוב הגפן 7? "
               "הצג כותרת נפרדת לכל מגרש ותחתיה תבליט לכל נתון.")


def test_a3_the_same_bullet_under_two_lot_headings_is_kept_and_no_lot_gets_the_others_datum_three_times(office, monkeypatch, tmp_path):
    doc = ingest_docx(office, monkeypatch, tmp_path, "R21_synthetic_two_lots.docx", LOTS_DOC)

    def lot_at(lines: list[tuple[int, str]], i: int, col: int) -> str | None:
        """The lot a place in line ``i`` at column ``col`` belongs to: the last lot named before it in its line, else the last one named in
        the nearest line above that names one."""
        for k in range(i, -1, -1):
            seg = lines[k][1][:col] if k == i else lines[k][1]
            found = [(m.end(), lot) for lot, pat in LOT_MENTION.items() for m in pat.finditer(seg)]
            if found:
                return max(found)[1]
        return None

    def placed(t: Turn21, anchor: re.Pattern) -> list[tuple[str | None, str]]:
        lines = t.body_lines
        return [(lot_at(lines, i, m.start()), ln.strip()) for i, (_, ln) in enumerate(lines) for m in anchor.finditer(ln)]

    def bullets_with_area(t: Turn21) -> list[str]:
        return [ln.strip() for _, ln in t.body_lines if BULLET.match(ln) and has(ln, LOT_AREA)]

    def both_lots_bullets_kept(t: Turn21) -> None:
        held = {lot for lot, _ in placed(t, LOT_AREA)}
        assert {"101", "102"} <= held, ("a lot's area bullet is gone", sorted(x for x in held if x), bullets_with_area(t), t.markdown[:500])
        assert len(bullets_with_area(t)) >= 2, ("fewer than two bullets state the area", bullets_with_area(t))
        # the binding record of the draft: the tidy-up and the repair may not take away a bullet the last verified draft held
        assert t.drafts, "no draft was captured"
        in_draft = sum(1 for ln in t.drafts[-1]["markdown"].splitlines() if BULLET.match(ln) and has(ln, LOT_AREA))
        assert len(bullets_with_area(t)) >= in_draft, ("the answer holds fewer area bullets than the last verified draft", in_draft,
                                                         len(bullets_with_area(t)))

    def values_stay_with_their_lot(t: Turn21) -> None:
        wrong = [(lot, ln[:200]) for own, pat in LOT_VALUE.items() for lot, ln in placed(t, pat) if lot not in (own, None)]
        assert not wrong, ("a lot's value stands under the other lot", wrong[:3])
        missing = [own for own, pat in LOT_VALUE.items() if not any(lot == own for lot, _ in placed(t, pat))]
        assert not missing, ("a lot's value does not stand under its own lot", missing, t.markdown[:500])

    def no_own_datum_removed_as_another_property(t: Turn21) -> None:
        bad = []
        for r in t.claim_rows:
            if r["action"] not in GONE_ACTIONS:
                continue
            row_text = r["text"] or ""
            for own, pat in LOT_VALUE.items():
                if has(row_text, LOT_MENTION[own]) and has(row_text, pat) and r["failure_kind"] == ANOTHER_PROPERTY:
                    bad.append((own, r["action"], r["check"], row_text[:160]))
        assert not bad, ("a lot's own datum was removed as another property's", bad[:3])
        words = [s["text"][:200] for s in t.spans if ANOTHER_PROPERTY_WORDS.search(s["text"])]
        assert not words, ("a server sentence says 'another property' about the lots of one property", words[:2])
        assert not REMOVAL_SENTENCE.search(t.markdown), ("the answer carries a removal sentence", t.markdown[:400])

    def body(o, c, cid, rec):
        t = ask(c, cid, A3_QUESTION, rec)
        reached = {"area_101": any(lot == "101" for lot, _ in placed(t, LOT_AREA)), "area_102": any(lot == "102" for lot, _ in placed(t, LOT_AREA)),
                   "value_101": stated(t, LOT_VALUE["101"]), "value_102": stated(t, LOT_VALUE["102"])}
        rec["info"] = base_info(t, reached, {"area": t.part_stage(LOT_AREA), "value_101": t.part_stage(LOT_VALUE["101"]),
                                             "value_102": t.part_stage(LOT_VALUE["102"])})
        rec["info"] |= {"shape_produced": len(bullets_with_area(t)) >= 2, "area_bullets": bullets_with_area(t),
                        "area_bullets_in_last_draft": sum(1 for ln in (t.drafts[-1]["markdown"].splitlines() if t.drafts else [])
                                                         if BULLET.match(ln) and has(ln, LOT_AREA)),
                        "area_places": placed(t, LOT_AREA), "value_places": {k: placed(t, p) for k, p in LOT_VALUE.items()},
                        "inline_spans": [(s["kind"], s["text"][:120], s["claim_ids"], s["datum"]) for s in t.inline_spans],
                        "removed_rows": [(r["action"], r["failure_kind"], r["check"], (r["text"] or "")[:120]) for r in t.claim_rows
                                         if r["action"] in GONE_ACTIONS],
                        "removed_as_another_property": t.removed_as_another_property}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "shape_produced", lambda: shape_not_produced(t, len(bullets_with_area(t)) >= 2, "no two bullets under lot headings state the area"))
        check(rec, "both_lots_bullets_kept", lambda: both_lots_bullets_kept(t))
        check(rec, "values_stay_with_their_lot", lambda: values_stay_with_their_lot(t))
        check(rec, "no_own_datum_removed_as_another_property", lambda: no_own_datum_removed_as_another_property(t))

    run_sequence("a3_lots_identical_bullet",
                 ("answered", "cited", "shape_produced", "both_lots_bullets_kept", "values_stay_with_their_lot",
                  "no_own_datum_removed_as_another_property"), office, body)


# --- a4: a follow-up question in the same conversation (AE8) ---------------------------------------------------------------------

A4_FIRST = "מהו השטח ומהו חוזה החכירה של קבוצת מגרשים א׳ בשכונת הרימון, ובאיזה תנאי כפוף התשלום לפי חוזה החכירה?"
A4_SECOND = "ומה השטח של קבוצת מגרשים ב׳?"


def test_a4_a_follow_up_answer_carries_no_condition_note_of_the_first_answer_three_times(office, monkeypatch):
    doc = ingest_pdf(office, monkeypatch, LOTS)

    def condition_in(t: Turn21) -> bool:
        return has(t.markdown, CONDITION)

    def no_carried_condition(first: Turn21, second: Turn21) -> None:
        bad = []
        carriers = [s for s in second.spans if CONDITION.search(s["text"]) or SERVER_CONDITION_NOTE.search(s["text"])]
        if carriers:
            bad.append(("a server span of the second answer carries the first answer's condition", [(s["kind"], s["text"][:160]) for s in carriers]))
        shared = {c for s in first.spans for c in s["claim_ids"]} & {c for s in second.spans for c in s["claim_ids"]}
        if shared:
            bad.append(("a claim id of the first answer's additions is in the second answer's registry", sorted(shared)))
        ledger = [(i["id"], sp) for i in second.condition_trace["items"] for sp in i["spans"] if has(sp[1] or "", CONDITION)]
        if ledger:
            bad.append(("an item of the second answer's ledger holds a condition span of the land authority", ledger[:3]))
        beside = [ln[:200] for ln in second.stated_lines if has(ln, AREA_B) and has(ln, CONDITION)]
        if beside:
            bad.append(("the condition stands beside the new group's datum", beside[:2]))
        assert not bad, bad

    def body(o, c, cid, rec):
        first = ask(c, cid, A4_FIRST, rec)
        rec["info"] = {"shape_produced": condition_in(first) and stated(first, AREA_A), "first_condition_in_answer": condition_in(first),
                       "first_spans": [(s["kind"], s["text"][:160], s["claim_ids"], s["datum"]) for s in first.spans]}
        second = ask(c, cid, A4_SECOND, rec)
        reached = {"area_b": stated(second, AREA_B)}
        rec["info"] |= base_info(second, reached, {"area_b": second.part_stage(AREA_B)}) | {
            "second_spans": [(s["kind"], s["text"][:160], s["claim_ids"], s["datum"]) for s in second.spans],
            "second_condition_anywhere": condition_in(second), "second_item_condition_spans": [
                (i["id"], i["spans"]) for i in second.condition_trace["items"] if i["spans"]]}
        check(rec, "answered_first", lambda: answered(first))
        check(rec, "shape_produced", lambda: shape_not_produced(first, condition_in(first) and stated(first, AREA_A),
                                                               "the first answer does not state group A's area with the payment condition"))
        check(rec, "answered_second", lambda: answered(second))
        check(rec, "cited_second", lambda: cites_document(second, doc))
        check(rec, "second_area_present", lambda: _all(reached, second))
        check(rec, "second_answer_carries_no_condition_of_the_first", lambda: no_carried_condition(first, second))

    run_sequence("a4_follow_up_no_carried_condition",
                 ("answered_first", "shape_produced", "answered_second", "cited_second", "second_area_present",
                  "second_answer_carries_no_condition_of_the_first"), office, body)


# --- a5: a reversed-stored left-to-right table (AE6) ------------------------------------------------------------------------------

def test_a5_a_reversed_stored_table_gives_the_value_under_the_right_column_label_three_times(office, monkeypatch):
    doc = ingest_pdf(office, monkeypatch, QUARTERS)
    reverse_stored_table(office, 0)  # test data only: the throwaway database, a table of a synthetic fixture (as AE6 builds it)

    def served(t: Turn21) -> list[dict]:
        return [s for s in t.sources if "Revenue" in (s.get("text") or "")]

    def labels_in(fragment: str) -> list[str]:
        return [k for k, pat in QUARTER_LABELS.items() if pat.search(fragment)]

    def value_under_the_right_label(t: Turn21) -> None:
        lines = t.stated_lines
        assert any(has(ln, QUARTER_VALUES["Q1"]) and "Q1" in labels_in(ln) for ln in lines), ("the first quarter's value is not given", t.markdown[:400])
        wrong = []
        for key, pat in QUARTER_VALUES.items():
            for ln in lines:
                if not has(ln, pat) or ln.lstrip().startswith("|"):
                    continue
                told = labels_in(ln)
                if len(told) == 1 and told[0] != key:
                    wrong.append((key, told[0], ln[:200]))
        assert not wrong, ("a quarter's value is told under another quarter's label", wrong[:3])
        # a table of the answer: each value stands in the column its header labels
        for table in _tables(lines):
            header = cells_of(table[0])
            for ln in table[1:]:
                cells = cells_of(ln)
                if len(cells) != len(header):
                    continue
                for k, cell in enumerate(cells):
                    for key, pat in QUARTER_VALUES.items():
                        if has(cell, pat) and (found := labels_in(header[k])) and found != [key]:
                            wrong.append((key, found, ln[:160]))
        assert not wrong, ("a quarter's value stands in another quarter's column", wrong[:3])

    def sources_in_drawn_order(t: Turn21) -> None:
        rows = served(t)
        assert rows, ("no cited source holds the table's row", [(s.get("id"), s.get("source_kind")) for s in t.sources])
        unpaired = [(s.get("id"), (s.get("text") or "")[:300]) for s in rows if not pairs_in_drawn_order(s["text"], "Revenue", QUARTER_PAIRS)]
        assert not unpaired, ("a cited source does not serve the row in drawn order with its column labels", unpaired[:2])
        reversed_row = [s.get("id") for s in rows if re.search(r"300\D+200\D+100", s["text"])]
        assert not reversed_row, ("a cited source serves the row reversed", reversed_row)

    def body(o, c, cid, rec):
        t = ask(c, cid, REVENUE_QUESTION, rec)
        reached = {"q1": stated(t, QUARTER_VALUES["Q1"])}
        rec["info"] = base_info(t, reached, {"q1": t.part_stage(QUARTER_VALUES["Q1"])})
        rec["info"] |= {"served_rows": [(s.get("id"), s.get("source_kind"), [ln[:200] for ln in (s.get("text") or "").splitlines() if "Revenue" in ln])
                                        for s in served(t)],
                        "claim_rows_of_q1": [(r["claim_id"], r["action"], r["verdict"], (r["text"] or "")[:120]) for r in t.claim_rows
                                             if has(r["text"] or "", QUARTER_VALUES["Q1"])]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "value_under_the_right_label", lambda: value_under_the_right_label(t))
        check(rec, "cited_sources_in_drawn_order", lambda: sources_in_drawn_order(t))
        check(rec, "value_not_removed", lambda: _no(verification_rows_with(t, QUARTER_VALUES["Q1"]), "the first quarter's value was removed by verification"))

    run_sequence("a5_reversed_stored_table",
                 ("answered", "cited", "value_under_the_right_label", "cited_sources_in_drawn_order", "value_not_removed"), office, body)


def _tables(lines: list[str]) -> list[list[str]]:
    """The markdown tables of the answer (runs of lines that start with a pipe), the separator lines left out."""
    out, cur = [], []
    for ln in lines:
        if ln.lstrip().startswith("|"):
            if not all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells_of(ln) if c):
                cur.append(ln)
        elif cur:
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


# --- the summary ----------------------------------------------------------------------------------------------------------

def build_summary() -> dict:
    turns = [t for r in RESULTS for t in r.get("turns") or []]
    purposes: dict[str, int] = {}
    statuses: dict[str, int] = {}
    for t in turns:
        for k, n in t["usage"]["purposes"].items():
            purposes[k] = purposes.get(k, 0) + n
        for k, n in t["usage"]["statuses"].items():
            statuses[k] = statuses.get(k, 0) + n
    costs = [t["usage"]["cost_usd"] for t in turns if t["usage"]["cost_usd"] is not None]
    names = list(dict.fromkeys(r["sequence"] for r in RESULTS))
    sequences = {}
    for name in names:
        s = sequence_summary(name)
        sequences[name] = {"passed": s["passed"], "passed_runs": s["passed_runs"], "shape_not_produced_runs": s["shape_not_produced_runs"],
                           "attempts": [{k: a[k] for k in ("attempt", "outcome", "checks", "failures", "stage", "info")} for a in s["attempts"]]}
    return {"model": MODEL, "turns": len(turns), "attempts": len(RESULTS), "sequences": sequences,
            "usage_purposes": purposes, "usage_statuses": statuses, "cost_usd": round(sum(costs), 4) if costs else None,
            "model_latency_s": round(sum(t["usage"]["model_latency_ms"] for t in turns) / 1000, 1),
            "wall_s": round(time.perf_counter() - STARTED, 1)}


def teardown_module(module) -> None:  # noqa: ARG001
    if RESULTS:
        flush()
        SUMMARY_FILE.write_text(json.dumps(build_summary(), ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        print(f"\nreal_model results: {os.path.relpath(RESULTS_FILE, ROOT)}\nsummary: {os.path.relpath(SUMMARY_FILE, ROOT)}")
