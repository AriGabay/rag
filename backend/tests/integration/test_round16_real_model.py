"""Round 16 acceptance with the real model (U7: R1-R6, R13; AE1, AE3): a part of a question that only a passage of text states
(no registered value, measurement or calculation behind it) reaches the answer. Four shapes (R13) on synthetic documents
generated here (python-docx, into a temporary directory; every name, address and figure invented) and ingested through the worker
into a synthetic office of the ``rag_test`` database; questions go through the chat API to the configured model and judge
(``gpt-6-luna`` with the repository's settings), whose key comes from the repository's ``.env`` and is never printed:

- a professional explanation: a standard states a rule and, in its own sentence, the reason for it;
- a condition: a rate and, in the next sentence, the condition under which it applies;
- a table: a table of rates by use, five rows;
- a comparison between two appraisals: the rate of each and its lease term, each to be attributed to its own appraisal.

Each shape is asked in a **new conversation, exactly three times**; every repetition is recorded - a failure in any repetition is
reported and a later pass never hides it - and a shape passes only when all three repetitions pass. Per repetition, by behaviour and
never by wording (the gating checks):

- the question is answered (not clarified, not failed);
- the part reached the answer: the words (or rows) that name it are in the answer;
- the part is attributed: the answer cites the passage's document, and, in the comparison, an item is cited only from its own
  appraisal's passage.

Informational (recorded, not gated): the completeness the turn reports, the stage the part reached the answer in (``first_pass``:
the model gave it; ``completed``: the server's completion round added it from the passage the turn had read; ``omitted``: it did
not reach the answer), the verification rounds' checks, the model calls and wall time. Per turn the whole record is written to
``real_documents/eval/round16/real-model/`` (gitignored and private: answers are kept for diagnosis), with a summary file next to
it listing every repetition of every shape.

Opt in with ``-m real_model``; skipped without a key. Host only, against ``rag_test`` (``tests/conftest.py`` guard)."""

from __future__ import annotations

import io
import json
import os
import re
import time
from datetime import UTC, datetime
from pathlib import Path

import pytest
from sqlalchemy import text

from app.config import get_settings
from app.db import tenant_tx
from tests.conftest import login
from tests.factories import make_office
from tests.integration.test_round10_real_model import new_conversation, upload, usage_of
from tests.integration.test_round10_real_model import real_key as real_key  # noqa: F401 - the key fixture

pytestmark = [pytest.mark.real_model, pytest.mark.db]

ROOT = Path(__file__).resolve().parents[3]
RESULTS_DIR = ROOT / "real_documents" / "eval" / "round16" / "real-model"
STAMP = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
RESULTS_FILE = RESULTS_DIR / f"results-{STAMP}.json"
SUMMARY_FILE = RESULTS_DIR / f"summary-{STAMP}.json"
RESULTS: list[dict] = []
ATTEMPTS = 3  # exactly three, each in a new conversation
STARTED = time.perf_counter()

ADMIN = "admin-r16@example.test"
MARK = "מסמך בדיקה סינתטי (סבב 16) - נתונים בדויים לבדיקת המערכת בלבד"
CITED = re.compile(r"\[([SV]\d+)\]")


# --- the synthetic documents -------------------------------------------------------------------------------------------

def docx_bytes(title: str, blocks: list[tuple[str, object]]) -> bytes:
    """A DOCX of ``blocks`` (``("h1", text)``, ``("p", text)`` or ``("table", [[cell, ...], ...])``) under ``title``, with Word
    heading styles so the reader finds the structure, and the synthetic mark as its first paragraph."""
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(MARK)
    for kind, body in blocks:
        if kind == "h1":
            doc.add_heading(str(body), 1)
        elif kind == "table":
            rows = list(body)  # type: ignore[arg-type]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, value in enumerate(cells):
                    table.cell(r, c).text = value
        else:
            doc.add_paragraph(str(body))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


STANDARD = {
    "title": "תקן סינתטי 77 - היוון הכנסה", "kind": "standard",
    "blocks": [("h1", "1. הגדרות"),
               ("p", "שיעור היוון הוא המקדם שבו ממירים הכנסה שנתית צפויה לשווי נכס."),
               ("h1", "2. בחירת שיעור ההיוון"),
               ("p", "על השמאי לקבוע את שיעור ההיוון לפי שיעור שוק."),
               ("p", "הנימוק לכך הוא ששיעור שוק משקף את התשואה שמשקיע מצפה לה מנכס חלופי בסיכון דומה, ולכן הוא מבטיח "
                     "שהשווי יהיה עקבי עם עסקאות בשוק."),
               ("h1", "3. דיווח"),
               ("p", "השמאי ירשום בחוות הדעת את שיעור ההיוון ואת מקורו.")]}
CONDITION = {
    "title": "שומה רחוב הצפצפה 4 עיר-סינתטית", "kind": "appraisal",
    "blocks": [("h1", "1. פרטי הנכס"), ("p", "דירת מגורים בת ארבעה חדרים ברחוב הצפצפה 4."),
               ("h1", "5. שיעור ההיוון"),
               ("p", "שיעור ההיוון שנקבע לדירות מגורים בשומה הוא 5.5%."),
               ("p", "שיעור זה חל רק בתנאי שהנכס מושכר במלואו ביום הקובע."),
               ("h1", "6. מסקנה"), ("p", "השווי נקבע לפי גישת ההשוואה.")]}
RATE_ROWS = [["משרדים", "6.5%"], ["מסחר", "7.25%"], ["חניה", "8.25%"], ["מחסנים", "9.5%"], ["מרפסות", "5.75%"]]
TABLE = {
    "title": "שומה רחוב השיטה 6 עיר-סינתטית", "kind": "appraisal",
    "blocks": [("h1", "1. פרטי הנכס"), ("p", "מתחם מעורב ברחוב השיטה 6."),
               ("h1", "4. שיעורי היוון לפי שימוש"), ("p", "שיעורי ההיוון שנקבעו לכל שימוש מפורטים בטבלה."),
               ("table", [["סוג שימוש", "שיעור היוון"], *RATE_ROWS])]}
LOTEM = {
    "title": "שומה רחוב הלוטם 2 עיר-סינתטית", "kind": "appraisal",
    "blocks": [("h1", "1. פרטי הנכס"), ("p", "משרד ברחוב הלוטם 2."),
               ("h1", "4. פרמטרי ההיוון וההשכרה"),
               ("p", "שיעור ההיוון בשומה הוא 6.25%."), ("p", "תקופת ההשכרה בשומה היא 3 שנים.")]}
NARKIS = {
    "title": "שומה רחוב הנרקיס 8 עיר-סינתטית", "kind": "appraisal",
    "blocks": [("h1", "1. פרטי הנכס"), ("p", "משרד ברחוב הנרקיס 8."),
               ("h1", "4. פרמטרי ההיוון וההשכרה"),
               ("p", "שיעור ההיוון בשומה הוא 7.5%."), ("p", "תקופת ההשכרה בשומה היא 5 שנים.")]}


# --- environment -------------------------------------------------------------------------------------------------------

@pytest.fixture
def office(db, client, monkeypatch, real_key):  # noqa: F811 - the key fixture is the imported one
    monkeypatch.setattr(get_settings(), "chat_run_inline", True)
    o = make_office(db, "משרד בדיקה סינתטי", ADMIN)
    with tenant_tx(o.ctx()) as conn:
        conn.execute(text("UPDATE office_settings SET cloud_llm_enabled = true"))
    login(client, ADMIN)
    o.client, o.ids = client, {}
    return o


def ingest(office, monkeypatch, tmp_path: Path, *specs: dict) -> None:
    """The synthetic documents ``specs`` generated and read by the worker; ``office.ids[title]`` is each document's id."""
    for spec in specs:
        path = tmp_path / (spec["title"].replace(" ", "_") + ".docx")
        path.write_bytes(docx_bytes(spec["title"], spec["blocks"]))
        office.ids[spec["title"]] = upload(office, path, spec["kind"], monkeypatch)


# --- turns and records -------------------------------------------------------------------------------------------------

def flush() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    RESULTS_FILE.write_text(json.dumps({"model": "gpt-6-luna (settings)", "attempts": RESULTS}, ensure_ascii=False, indent=2,
                                       default=str), encoding="utf-8")


class Turn:
    """A finished turn: the message and answer, the checks of each verification round (its diagnostics) and the usage."""

    def __init__(self, c, m: dict, question: str, wall: float) -> None:
        self.m, self.question = m, question
        self.answer = m.get("answer") or {}
        d = c.get(f"/api/chat/messages/{m['id']}/diagnostics")
        body = d.json() if d.status_code == 200 else {}
        self.rounds = [[{"check": p.get("check"), "kind": p.get("kind"), "component": p.get("component")} for p in r]
                       for r in body.get("rounds") or []]
        self.tools = [t.get("tool") for t in (body.get("timing") or {}).get("tools") or []]
        self.usage = usage_of(m, wall)

    @property
    def markdown(self) -> str:
        return self.answer.get("markdown") or ""

    @property
    def sources(self) -> list[dict]:
        return self.answer.get("sources") or []

    @property
    def completeness(self) -> dict:
        return (self.answer.get("verification") or {}).get("completeness") or {}

    @property
    def completed_by_server(self) -> bool:
        """Whether a verification round raised a completion from read text (``REQ_TEXT_FOUND``, ``check="text_found"``)."""
        return any(p["check"] == "text_found" for r in self.rounds for p in r)

    def lines_with(self, needle: str) -> list[str]:
        return [ln for ln in self.markdown.splitlines() if needle in ln]

    def documents_cited_in(self, line: str) -> set[str]:
        """The documents of the sources the line cites."""
        by_id = {s.get("id"): s.get("document_id") for s in self.sources}
        return {by_id[i] for i in CITED.findall(line) if i in by_id}

    def record(self) -> dict:
        return {"question": self.question, "message_id": self.m.get("id"), "message_status": self.m.get("status"),
                "answer_status": self.answer.get("status"), "completeness": self.completeness,
                "completed_by_server": self.completed_by_server, "rounds": self.rounds, "tools": self.tools,
                "cited": [{"id": s.get("id"), "document_id": s.get("document_id"), "title": s.get("title"),
                           "location": s.get("location"), "source_kind": s.get("source_kind")} for s in self.sources],
                "usage": self.usage, "markdown": self.markdown}


def ask(c, cid: str, question: str, rec: dict) -> Turn:
    """One turn through the API; recorded (and flushed) before anything is asserted about it."""
    started = time.perf_counter()
    r = c.post(f"/api/chat/conversations/{cid}/messages", json={"content": question})
    assert r.status_code == 200, r.text
    m = c.get(f"/api/chat/messages/{r.json()['assistant']['id']}").json()
    t = Turn(c, m, question, time.perf_counter() - started)
    rec["turns"].append(t.record())
    flush()
    assert m["status"] == "done", (m.get("status"), m.get("error"))
    return t


def check(rec: dict, name: str, fn) -> None:
    """Run one named assertion of a repetition and record its outcome; a failure never stops the repetition's other checks."""
    try:
        fn()
    except AssertionError as exc:
        rec["checks"][name] = {"outcome": "fail", "failure": str(exc)[:2000]}
    else:
        rec["checks"][name] = {"outcome": "pass"}
    flush()


def run_sequence(sequence: str, checks: tuple[str, ...], office, body) -> None:
    """``body(office, client, conversation_id, record)`` in a new conversation, exactly ``ATTEMPTS`` times; every repetition is
    recorded and none is skipped after a failure. The test fails when any check of any repetition failed (or never ran because
    a turn did not finish), with every failure in the message."""
    failures: list[str] = []
    for n in range(1, ATTEMPTS + 1):
        rec = {"sequence": sequence, "attempt": n, "turns": [], "checks": {}, "info": {}, "outcome": None}
        RESULTS.append(rec)
        try:
            body(office, office.client, new_conversation(office.client), rec)
        except AssertionError as exc:
            rec["aborted"] = str(exc)[:2000]
        for name in checks:
            rec["checks"].setdefault(name, {"outcome": "fail", "failure": "never ran: " + rec.get("aborted", "?")})
        rec["outcome"] = "pass" if all(v["outcome"] == "pass" for v in rec["checks"].values()) else "fail"
        flush()
        failures += [f"attempt {n} / {name}: {v['failure']}" for name, v in rec["checks"].items() if v["outcome"] == "fail"]
    assert not failures, "\n".join(failures)


def stage(t: Turn, reached: bool) -> str:
    """The stage the part reached the answer in: ``first_pass`` (the model gave it), ``completed`` (the server's completion round
    added it) or ``omitted``."""
    return "omitted" if not reached else "completed" if t.completed_by_server else "first_pass"


def answered(t: Turn) -> None:
    assert t.answer.get("status") != "clarification", ("a clarification was asked", t.markdown[:300])
    assert t.markdown.strip(), "an empty answer"


def cites_document(t: Turn, document: str) -> None:
    assert any(s.get("document_id") == document for s in t.sources), (
        "the passage's document is not cited", [(s.get("id"), s.get("title")) for s in t.sources])


def one_shape(sequence: str, question: str, parts: dict[str, tuple[str, ...]], documents: list[str], office, extra=None) -> None:
    """The common sequence: ``question`` asked in a new conversation three times; the gating checks are ``answered``, ``part``
    (every group of ``parts``, a name -> words that must all be in the answer) and ``cited`` (each of ``documents``, a title,
    cited); ``extra(office, turn)`` adds more. Informational: the stage of every part, the completeness, the rounds."""
    names = ("answered", "part", "cited", *(("attributed",) if extra else ()))

    def body(o, c, cid, rec):
        t = ask(c, cid, question, rec)
        text_ = re.sub(r"\s+", " ", t.markdown)
        reached = {name: all(w in text_ for w in words) for name, words in parts.items()}
        rec["info"] = {"parts_reached": reached, "stage": {n: stage(t, r) for n, r in reached.items()},
                       "completeness": t.completeness.get("status"), "completed_by_server": t.completed_by_server,
                       "tools": t.tools}
        check(rec, "answered", lambda: answered(t))
        check(rec, "part", lambda: _all_reached(reached))
        check(rec, "cited", lambda: [cites_document(t, o.ids[title]) for title in documents])
        if extra:
            check(rec, "attributed", lambda: extra(o, t))

    run_sequence(sequence, names, office, body)


def _all_reached(reached: dict[str, bool]) -> None:
    missing = [n for n, ok in reached.items() if not ok]
    assert not missing, ("the part did not reach the answer", missing)


# =========================================================================================================================
# the four shapes
# =========================================================================================================================

def test_s1_a_professional_explanation_the_standard_states_reaches_the_answer_three_times(office, monkeypatch, tmp_path):
    ingest(office, monkeypatch, tmp_path, STANDARD)
    one_shape("s1_professional_explanation", "לפי תקן סינתטי 77, איך קובעים את שיעור ההיוון ומה הנימוק לכך?",
              {"rule": ("שיעור שוק",), "reason": ("חלופי", "סיכון דומה")}, [STANDARD["title"]], office)


def test_s2_a_condition_of_a_rate_reaches_the_answer_three_times(office, monkeypatch, tmp_path):
    ingest(office, monkeypatch, tmp_path, CONDITION)
    one_shape("s2_condition", "מהו שיעור ההיוון שנקבע בשומה לדירות מגורים, ובאיזה תנאי הוא חל?",
              {"rate": ("5.5",), "condition": ("מושכר במלואו",)}, [CONDITION["title"]], office)


def test_s3_every_row_of_a_table_reaches_the_answer_three_times(office, monkeypatch, tmp_path):
    ingest(office, monkeypatch, tmp_path, TABLE)
    one_shape("s3_table_rows", "מהם שיעורי ההיוון שנקבעו בשומה לפי סוג שימוש?",
              {label: (label, value) for label, value in RATE_ROWS}, [TABLE["title"]], office)


def test_s4_each_appraisals_item_reaches_the_answer_with_its_own_attribution_three_times(office, monkeypatch, tmp_path):
    ingest(office, monkeypatch, tmp_path, LOTEM, NARKIS)

    def attributed(o, t: Turn) -> None:
        own = {"6.25": o.ids[LOTEM["title"]], "3 שנים": o.ids[LOTEM["title"]],
               "7.5": o.ids[NARKIS["title"]], "5 שנים": o.ids[NARKIS["title"]]}
        wrong = []
        for value, doc in own.items():
            cited = [t.documents_cited_in(ln) for ln in t.lines_with(value)]
            if not any(doc in c for c in cited) or any(c and doc not in c for c in cited):
                wrong.append((value, [sorted(c) for c in cited]))
        assert not wrong, ("an item is not attributed to its own appraisal", wrong)

    one_shape("s4_comparison_two_appraisals",
              "השווה בין השומה של רחוב הלוטם 2 לשומה של רחוב הנרקיס 8: מה שיעור ההיוון ומה תקופת ההשכרה בכל אחת?",
              {"lotem_rate": ("6.25",), "lotem_term": ("3 שנים",), "narkis_rate": ("7.5",), "narkis_term": ("5 שנים",)},
              [LOTEM["title"], NARKIS["title"]], office, extra=attributed)


# --- the summary -------------------------------------------------------------------------------------------------------

def build_summary() -> dict:
    turns = [t for r in RESULTS for t in r.get("turns") or []]
    purposes: dict[str, int] = {}
    statuses: dict[str, int] = {}
    for t in turns:
        for k, n in t["usage"]["purposes"].items():
            purposes[k] = purposes.get(k, 0) + n
        for k, n in t["usage"]["statuses"].items():
            statuses[k] = statuses.get(k, 0) + n
    costs = [t["usage"]["cost_usd"] for t in turns if t["usage"]["cost_usd"] is not None]
    sequences: dict[str, list[dict]] = {}
    for r in RESULTS:
        sequences.setdefault(r["sequence"], []).append(
            {"attempt": r["attempt"], "outcome": r["outcome"], "checks": {k: v["outcome"] for k, v in r["checks"].items()},
             "failures": {k: v["failure"] for k, v in r["checks"].items() if v["outcome"] == "fail"},
             "parts_reached": (r.get("info") or {}).get("parts_reached"), "stage": (r.get("info") or {}).get("stage"),
             "completeness": (r.get("info") or {}).get("completeness")})
    return {"model": "gpt-6-luna (settings)", "turns": len(turns), "attempts": len(RESULTS),
            "sequences": {k: {"attempts": v, "passed": len(v) == ATTEMPTS and all(a["outcome"] == "pass" for a in v),
                              "passed_runs": f"{sum(a['outcome'] == 'pass' for a in v)}/{len(v)}"}
                          for k, v in sequences.items()},
            "usage_purposes": purposes, "usage_statuses": statuses, "cost_usd": round(sum(costs), 4) if costs else None,
            "model_latency_s": round(sum(t["usage"]["model_latency_ms"] for t in turns) / 1000, 1),
            "wall_s": round(time.perf_counter() - STARTED, 1)}


def teardown_module(module) -> None:  # noqa: ARG001
    if RESULTS:
        flush()
        SUMMARY_FILE.write_text(json.dumps(build_summary(), ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nreal_model results: {os.path.relpath(RESULTS_FILE, ROOT)}\nsummary: {os.path.relpath(SUMMARY_FILE, ROOT)}")
