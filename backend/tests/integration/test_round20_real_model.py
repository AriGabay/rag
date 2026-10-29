"""Round 20 acceptance with the real model (U6: R20; AE1, AE6, AE7, AE8 and the cell-level final check of a table row): the shapes the
round-20 fixes handle, shown on the configured model and judge (``gpt-6-luna`` with the repository's settings, no override; the key
comes from the repository's ``.env`` through the round-10 ``real_key`` fixture and is never printed). Synthetic documents only (every
name, address and figure invented): the PDFs of ``tests/fixtures/round20`` (``scripts/generate_fixtures.py --only round20``; the
contents are the manifest's), ingested through the worker into a synthetic office of the ``rag_test`` database. The one exception
is sequence ``e1``: no round-20 fixture holds a payments paragraph, so its DOCX (two payments, each with a value and a date, in one
sentence - the text of the unit tests' ``PAYMENTS``) is written into the test's own temporary directory, never into the repository.

GATED (all three repetitions must pass):

- ``e1`` two payments in one quoted sentence (AE1): asked for the amount and the date of the two payments, both values and both
  dates reach the answer, each value stands with its own date (never with the other payment's date), and no removal sentence stands;
- ``e2`` a table row with a note cell (the cell-level final check): R20b (lot groups and their lease contracts) asked as a table
  with columns the document answers (the group, the area, the contract) and columns it only partly answers (the payment condition,
  which only group A carries, and the end date of the lease, which it does not hold). The shape is "a table whose row of a group
  holds a note cell"; a repetition in which the model writes no such shape is recorded as ``shape not produced`` and counted
  separately (it is not a pass). When the shape is produced: the supported cells of both rows (the area and the contract) stay in
  their cells, the table keeps one column count in every line, and the unsupported cell carries no invented figure or date and not
  the other group's footnote. How the note cell ended (``not_stated`` by the model, ``marked_by_server``, ``cut_empty``,
  ``invented_let_through`` or ``text``) is recorded, not forced.

REPORTED (the record is the result; only a finished turn is asserted, except where stated):

- ``e3`` plan A's column without plan B's figures (R20a): asked for the rights of both plans, each in a column of its own, no plan-B
  figure stands under plan A (asserted), the figures of each plan are recorded with the side they stand on;
- ``e4`` a footnote condition on its own lot group only (R20b): where the condition (approval of the land authority) stands -
  with group A, with group B, or with neither - recorded;
- ``e5`` a name/address unit (R20e): asked for the value of the property named in its title by name; the value is kept, it is not
  removed as another property's, the other property's value does not stand - recorded;
- ``e6`` an injected wrong-plan value: skipped, with the reason below.

Each sequence is asked in a **new conversation, exactly three times**; every repetition is recorded before anything is asserted - a
failure in any repetition is reported and a later pass never hides it - and a gated sequence passes only when all three
repetitions pass. The checks are by behaviour and never by wording (a pattern part is matched, so a Hebrew definite article or
prefix does not change the outcome).

Recorded per repetition (``real_documents/eval/round20/real-model/``, gitignored: the documents are synthetic, so the answers are
kept for diagnosis): the whole turn (final text, rounds, tools, item ledger, the request components, the spans the server wrote
with their kind and text, the removals, usage), the pass or fail of each check, and the stage that produced the part -
``agent_draft``, ``completion_repair``, ``server_attachment`` or ``omitted``. A summary file per sequence is written when the
sequence ends, and one for the run at the end of the module.

Opt in with ``-m real_model``; skipped without a key. Host only, against ``rag_test`` (``tests/conftest.py`` guard)."""

from __future__ import annotations

import io
import json
import os
import re
import time
from datetime import UTC, datetime
from pathlib import Path

import pytest
from sqlalchemy import text

from app.config import get_settings
from app.db import tenant_tx
from tests.conftest import login
from tests.factories import make_office
from tests.integration.test_round10_real_model import new_conversation, upload
from tests.integration.test_round10_real_model import real_key as real_key  # noqa: F401 - the key fixture
from tests.integration.test_round17_real_model import CITED, answered, cites_document, has
from tests.integration.test_round18_real_model import figure
from tests.integration.test_round19_real_model import (
    REMOVAL_SENTENCE,
    Turn19,
    base_info,
    cells_of,
    no_removal_as_another_property,
    stated,
)

pytestmark = [pytest.mark.real_model, pytest.mark.db]

ROOT = Path(__file__).resolve().parents[3]
FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
MANIFEST = json.loads((FIXTURES / "round20" / "manifest.json").read_text(encoding="utf-8"))
DOCS = MANIFEST["documents"]
SYNTHETIC_MARKER = MANIFEST["synthetic_marker"]
PLANS, LOTS, NAMES = DOCS["two_plans"], DOCS["lot_groups"], DOCS["name_address"]
RESULTS_DIR = ROOT / "real_documents" / "eval" / "round20" / "real-model"
STAMP = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
RESULTS_FILE = RESULTS_DIR / f"results-{STAMP}.json"
SUMMARY_FILE = RESULTS_DIR / f"summary-{STAMP}.json"
RESULTS: list[dict] = []
ATTEMPTS = 3  # exactly three, each in a new conversation
STARTED = time.perf_counter()
MODEL = "gpt-6-luna (settings)"

ADMIN = "admin-r20@example.test"

# the server's marks beside a part it cut or did not verify (the cell-level mark of a table row, a removal note) and the notes it keeps
SERVER_MARK = re.compile(r"הושמט|לא\s+אומת|לא\s+ניתן\s+לאמת|לא\s+נבדק")
FLAG_KINDS = {"partial_mark", "qualifier", "removal_note", "replacement", "marker", "limit_sentence", "cell_mark"}
NOT_STATED = re.compile(r"לא\s+(?:צוין|צוינה|נמסר|נקבע|מצוין|ידוע|קיים|פורט|נאמר|נמצא|הוגדר|מופיע)|אין|ללא|—\s*$|^[-–—־\s]*$")


# --- the synthetic payments document (the one document written by the test, in a temporary directory) --------------------------

PAYMENTS_TITLE = "שומת הדגמה לנכס ברחוב הצבר 6 — לוח תשלומים"
PAYMENTS_PARAGRAPH = ("התשלום הראשון בסך 400,000 ₪ ישולם ביום 01.03.2027 והתשלום השני בסך 650,000 ₪ ישולם ביום 01.09.2027. "
                      "יתרת התמורה תשולם במועד המסירה.")
FIRST_VALUE, SECOND_VALUE = figure("400000"), figure("650000")
# a date written 01.03.2027, 1/3/2027 or "1 במרץ 2027" (the second: 09 / ספטמבר), never as part of a longer number
FIRST_DATE = re.compile(r"(?<![\d.])0?1\s*[./\-]\s*0?3\s*[./\-]\s*(?:20)?2027(?!\d)|(?<!\d)1\s+במרץ\s+2027|(?<!\d)מרץ\s+2027")
SECOND_DATE = re.compile(r"(?<![\d.])0?1\s*[./\-]\s*0?9\s*[./\-]\s*(?:20)?2027(?!\d)|(?<!\d)1\s+בספטמבר\s+2027|(?<!\d)ספטמבר\s+2027")


def payments_docx() -> bytes:
    """A DOCX with Word heading styles: a purpose paragraph, then the payment schedule in one sentence (two payments, each with a
    value and a date) and a closing sentence about the balance."""
    from docx import Document

    doc = Document()
    doc.add_heading(PAYMENTS_TITLE, 0)
    doc.add_paragraph(f"{SYNTHETIC_MARKER} — כל הנתונים בדויים")
    doc.add_heading("1. מטרת השומה", 1)
    doc.add_paragraph("שומה זו נערכה לצורך הדגמה בלבד (מסמך סינתטי).")
    doc.add_heading("3. לוח התשלומים", 1)
    doc.add_paragraph(PAYMENTS_PARAGRAPH)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- what each sequence looks for -----------------------------------------------------------------------------------------

# e2/e4: the two lot groups of R20b by the area that tells them apart, and the footnote's condition (the land authority's approval)
(GROUP_A_ROW, GROUP_B_ROW) = LOTS["rows"]
AREA_A, AREA_B = figure(GROUP_A_ROW[2].replace(",", "")), figure(GROUP_B_ROW[2].replace(",", ""))
GROUPS = {"A": re.compile(r"(?<![\w])א[׳'’]|ראשון"), "B": re.compile(r"(?<![\w])ב[׳'’]|שני")}
CONDITION = re.compile(r"מינהל|מנהל\s+מקרקעי|רמ״י|מקרקעי\s+ישראל")
CONTRACT_CELL = re.compile(r"^(?:ה?חוזה\s+)?(?:ה?חכירה\s+)?(?:ה)?(?:ראשון|שני)$")
SEPARATOR_CELL = re.compile(r"^:?-{2,}:?$")
# e3: the two plans and the figures of each
PLAN_WORDS = {"A": re.compile(r"11111"), "B": re.compile(r"22222")}
PLAN_FIGURES = {"A": {"main": figure("1200"), "service": figure("300")},
                "B": {"main": figure("1800"), "service": figure("450")}}
# e5: the value of the asked property and the other property's
NAME_VALUE, OTHER_NAME_VALUE = figure(NAMES["reports"][0]["value"].replace(",", "")), figure(NAMES["reports"][1]["value"].replace(",", ""))


# --- environment ----------------------------------------------------------------------------------------------------------

@pytest.fixture
def office(db, client, monkeypatch, real_key):  # noqa: F811 - the key fixture is the imported one
    monkeypatch.setattr(get_settings(), "chat_run_inline", True)
    o = make_office(db, "משרד בדיקה סינתטי", ADMIN)
    with tenant_tx(o.ctx()) as conn:
        conn.execute(text("UPDATE office_settings SET cloud_llm_enabled = true"))
    login(client, ADMIN)
    o.client, o.ids = client, {}
    return o


def ingest_pdf(office, monkeypatch, spec: dict) -> str:
    """A round-20 PDF fixture read by the worker; ``office.ids[...]`` is its document id."""
    key = spec.get("file")
    office.ids[key] = upload(office, FIXTURES / spec["file"], "appraisal", monkeypatch)
    return office.ids[key]


def ingest_payments(office, monkeypatch, tmp_path: Path) -> str:
    path = tmp_path / "R20_synthetic_payments_schedule.docx"
    path.write_bytes(payments_docx())
    office.ids["payments"] = upload(office, path, "appraisal", monkeypatch)
    return office.ids["payments"]


# --- turns and records ----------------------------------------------------------------------------------------------------

def flush() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    RESULTS_FILE.write_text(json.dumps({"model": MODEL, "attempts": RESULTS}, ensure_ascii=False, indent=2, default=str),
                            encoding="utf-8")


def ask(c, cid: str, question: str, rec: dict) -> Turn19:
    """One turn through the API; recorded (and flushed) before anything is asserted about it."""
    started = time.perf_counter()
    r = c.post(f"/api/chat/conversations/{cid}/messages", json={"content": question})
    assert r.status_code == 200, r.text
    m = c.get(f"/api/chat/messages/{r.json()['assistant']['id']}").json()
    t = Turn19(c, m, question, time.perf_counter() - started)
    rec["turns"].append(t.record())
    flush()
    assert m["status"] == "done", (m.get("status"), m.get("error"))
    return t


def check(rec: dict, name: str, fn) -> None:
    """Run one named assertion of a repetition and record its outcome; a failure never stops the repetition's other checks."""
    try:
        fn()
    except AssertionError as exc:
        rec["checks"][name] = {"outcome": "fail", "failure": str(exc)[:2000]}
    else:
        rec["checks"][name] = {"outcome": "pass"}
    flush()


def sequence_summary(sequence: str) -> dict:
    runs = [r for r in RESULTS if r["sequence"] == sequence]
    attempts = []
    for r in runs:
        info = r.get("info") or {}
        turn = (r.get("turns") or [{}])[-1]
        attempts.append({"attempt": r["attempt"], "outcome": r["outcome"],
                         "checks": {k: v["outcome"] for k, v in r["checks"].items()},
                         "failures": {k: v["failure"] for k, v in r["checks"].items() if v["outcome"] == "fail"},
                         "stage": info.get("stage"), "info": {k: v for k, v in info.items() if k != "stage"},
                         "final_text": turn.get("markdown"), "server_spans": turn.get("server_spans"),
                         "rounds": turn.get("rounds"), "tools": turn.get("tools")})
    passed = sum(a["outcome"] == "pass" for a in attempts)
    not_produced = sum(1 for a in attempts if a["info"].get("shape_produced") is False)
    return {"model": MODEL, "sequence": sequence, "attempts": attempts, "passed_runs": f"{passed}/{len(attempts)}",
            "shape_not_produced_runs": not_produced, "passed": len(attempts) == ATTEMPTS and passed == ATTEMPTS}


def run_sequence(sequence: str, checks: tuple[str, ...], office, body, asserted: tuple[str, ...] | None = None) -> None:
    """``body(office, client, conversation_id, record)`` in a new conversation, exactly ``ATTEMPTS`` times; every repetition is
    recorded and none is skipped after a failure. The summary file of the sequence is written before anything is asserted. The
    test fails when a check in ``asserted`` (default: every check) failed in any repetition (or never ran because a turn did not
    finish), with every failure in the message; the other checks of a reported sequence are recorded and never fail the test."""
    failures: list[str] = []
    for n in range(1, ATTEMPTS + 1):
        rec = {"sequence": sequence, "attempt": n, "turns": [], "checks": {}, "info": {}, "outcome": None}
        RESULTS.append(rec)
        try:
            body(office, office.client, new_conversation(office.client), rec)
        except AssertionError as exc:
            rec["aborted"] = str(exc)[:2000]
        for name in checks:
            rec["checks"].setdefault(name, {"outcome": "fail", "failure": "never ran: " + rec.get("aborted", "?")})
        rec["outcome"] = "pass" if all(v["outcome"] == "pass" for v in rec["checks"].values()) else "fail"
        flush()
        failures += [f"attempt {n} / {name}: {v['failure']}" for name, v in rec["checks"].items()
                     if v["outcome"] == "fail" and (asserted is None or name in asserted)]
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    (RESULTS_DIR / f"summary-{sequence}-{STAMP}.json").write_text(
        json.dumps(sequence_summary(sequence), ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    assert not failures, "\n".join(failures)


# --- shared readers -------------------------------------------------------------------------------------------------------

def sides_in(fragment: str, words: dict[str, re.Pattern]) -> set[str]:
    return {k for k, w in words.items() if w.search(fragment)}


def placements(t: Turn19, anchor: re.Pattern, words: dict[str, re.Pattern]) -> list[tuple[str, str]]:
    """For each place the answer states ``anchor``: (the side the place tells - a key of ``words`` or ``unknown`` - and the line). A
    table cell is told by the header of its column, else by the label of its row; a sentence is told by the side it names, else by
    the nearest line before it that names one (three lines back at most); a place that names several sides is not told."""
    lines, out = t.stated_lines, []
    for i, ln in enumerate(lines):
        if not has(ln, anchor):
            continue
        told: set[str] = set()
        if ln.lstrip().startswith("|"):
            cells = cells_of(ln)
            idx = next((k for k, c in enumerate(cells) if has(c, anchor)), 0)
            start = i
            while start > 0 and lines[start - 1].lstrip().startswith("|"):
                start -= 1
            header = cells_of(lines[start]) if start != i else []
            if idx > 0 and len(header) == len(cells):
                told = sides_in(header[idx], words)
            if not told and idx > 0:
                told = sides_in(cells[0], words)
        else:
            told = sides_in(ln, words)
            for back in range(1, 4):
                if told or i - back < 0:
                    break
                told = sides_in(lines[i - back], words)
        out.append((next(iter(told)) if len(told) == 1 else "unknown", ln))
    return out


def tables_of(t: Turn19) -> list[list[str]]:
    """The markdown tables of the answer (runs of lines that start with a pipe), the server's gap lines left out."""
    out, cur = [], []
    for ln in t.stated_lines:
        if ln.lstrip().startswith("|"):
            cur.append(ln)
        elif cur:
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


def is_separator(line: str) -> bool:
    cells = cells_of(line)
    return bool(cells) and all(SEPARATOR_CELL.match(c.replace(" ", "")) for c in cells if c)


# --- e1: two payments in one quote ----------------------------------------------------------------------------------------

PAYMENT_PARTS = {"first_value": FIRST_VALUE, "second_value": SECOND_VALUE, "first_date": FIRST_DATE, "second_date": SECOND_DATE}


def wrong_pairing(t: Turn19) -> list[str]:
    """Lines that join a payment's value with the other payment's date: the first value with the second date and no first date, the
    second value with the first date and no second date, or - in a line with both values and both dates - the two in different
    orders."""
    bad = []
    for ln in t.stated_lines:
        v1, v2, d1, d2 = (pat.search(ln) for pat in (FIRST_VALUE, SECOND_VALUE, FIRST_DATE, SECOND_DATE))
        if (v1 and d2 and not d1 and not v2) or (v2 and d1 and not d2 and not v1):
            bad.append(ln[:240])
        elif v1 and v2 and d1 and d2 and (v1.start() < v2.start()) != (d1.start() < d2.start()):
            bad.append(ln[:240])
    return bad


def test_e1_two_payments_in_one_quote_are_both_kept_with_their_values_and_dates_three_times(office, monkeypatch, tmp_path):
    doc = ingest_payments(office, monkeypatch, tmp_path)

    def each_with_its_date(t: Turn19) -> None:
        lines = t.stated_lines
        missing = []
        for name, value, date in (("first", FIRST_VALUE, FIRST_DATE), ("second", SECOND_VALUE, SECOND_DATE)):
            same_line = any(has(ln, value) and has(ln, date) for ln in lines)
            near = any(has(" ".join(lines[max(0, i - 1):i + 2]), date) for i, ln in enumerate(lines) if has(ln, value))
            if not (same_line or near):
                missing.append(name)
        assert not missing, ("a payment is stated without its own date next to its value", missing, t.markdown[:500])

    def body(o, c, cid, rec):
        t = ask(c, cid, "מהם הסכום והמועד של כל אחד משני התשלומים הראשונים בנכס ברחוב הצבר 6?", rec)
        reached = {n: stated(t, pat) for n, pat in PAYMENT_PARTS.items()}
        same_line = {n: any(has(ln, v) and has(ln, d) for ln in t.stated_lines)
                     for n, v, d in (("first", FIRST_VALUE, FIRST_DATE), ("second", SECOND_VALUE, SECOND_DATE))}
        rec["info"] = base_info(t, reached, {n: t.part_stage(p) for n, p in PAYMENT_PARTS.items()})
        rec["info"] |= {"value_and_date_on_one_line": same_line, "wrong_pairings": wrong_pairing(t),
                        "items_included": sum(i["state"] == "included" for i in t.items),
                        "removed_as_another_property": t.removed_as_another_property}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "both_values_kept", lambda: _all({k: v for k, v in reached.items() if k.endswith("value")}, t))
        check(rec, "both_dates_kept", lambda: _all({k: v for k, v in reached.items() if k.endswith("date")}, t))
        check(rec, "each_value_with_its_date", lambda: each_with_its_date(t))
        check(rec, "no_value_with_the_other_payments_date", lambda: _no(wrong_pairing(t), "a value stands with the other payment's date"))
        check(rec, "no_removal_sentence", lambda: _no([t.markdown] if REMOVAL_SENTENCE.search(t.markdown) else [],
                                                       "the answer carries a removal sentence"))

    run_sequence("e1_two_payments_in_one_quote",
                 ("answered", "cited", "both_values_kept", "both_dates_kept", "each_value_with_its_date",
                  "no_value_with_the_other_payments_date", "no_removal_sentence"), office, body)


def _all(reached: dict[str, bool], t: Turn19) -> None:
    missing = [n for n, ok in reached.items() if not ok]
    assert not missing, ("a part did not reach the answer", missing, t.markdown[:500])


def _no(found: list, message: str) -> None:
    assert not found, (message, found[:3])


# --- e2: a table row with a note cell -------------------------------------------------------------------------------------

E2_QUESTION = ("הצג בטבלה את שתי קבוצות המגרשים בשכונת הרימון: לכל קבוצה את שמה, השטח שלה במ״ר, חוזה החכירה שלה, "
               "ההערה על תנאי התשלום לפי החוזה ותאריך סיום חוזה החכירה.")


def core_cell(idx: int, cell: str) -> bool:
    """A cell the document supports for sure: the first cell (the group), the area, or the name of the contract."""
    bare = CITED.sub("", cell).strip()
    return idx == 0 or has(bare, AREA_A) or has(bare, AREA_B) or bool(CONTRACT_CELL.match(re.sub(r"\s+", " ", bare)))


def row_of(table: list[str], area: re.Pattern) -> list[str] | None:
    for ln in table:
        if not is_separator(ln) and has(ln, area):
            return cells_of(ln)
    return None


def note_cells(cells: list[str]) -> list[str]:
    return [c for i, c in enumerate(cells) if not core_cell(i, c)]


def note_state(cells: list[str], t: Turn19) -> str:
    """How the note cells of a row ended: ``cut_empty`` (nothing, a dash), ``marked_by_server`` (the server's mark or a span the
    server wrote beside it), ``invented_let_through`` (a figure or a date the document does not hold), ``not_stated`` (the model says
    the document does not give it) or ``text``."""
    notes = [CITED.sub("", c).strip() for c in note_cells(cells)]
    joined = " ".join(notes)
    if any(SERVER_MARK.search(n) for n in notes) or any(s["kind"] in FLAG_KINDS for s in t.spans):
        return "marked_by_server"
    if re.search(r"\d", re.sub(r"\s+", "", joined)):
        return "invented_let_through"
    if notes and all(re.fullmatch(r"[-–—־\s]*", n) for n in notes):
        return "cut_empty"
    if notes and any(NOT_STATED.search(n) for n in notes):
        return "not_stated"
    return "text"


def test_e2_a_table_row_keeps_its_supported_cells_and_cuts_or_marks_the_unsupported_cell_three_times(office, monkeypatch):
    doc = ingest_pdf(office, monkeypatch, LOTS)

    def find(t: Turn19) -> dict:
        """The table that holds both groups' rows (by their areas) and the shape facts."""
        for table in tables_of(t):
            a, b = row_of(table, AREA_A), row_of(table, AREA_B)
            if a is not None and b is not None:
                header = cells_of(table[0])
                return {"table": table, "row_a": a, "row_b": b, "header": header, "notes_a": note_cells(a), "notes_b": note_cells(b)}
        return {}

    def shape_produced(t: Turn19) -> None:
        f = find(t)
        assert f and (f["notes_a"] or f["notes_b"]), ("SHAPE NOT PRODUCED: no table row of a lot group holds a note cell",
                                                      [tb[:4] for tb in tables_of(t)][:2], t.markdown[:500])

    def supported_cells_intact(t: Turn19) -> None:
        f = find(t)
        assert f, ("the rows of the two groups are not in one table", t.markdown[:500])
        bad = []
        for name, row, word in (("A", f["row_a"], "ראשון"), ("B", f["row_b"], "שני")):
            if not any(has(c, AREA_A if name == "A" else AREA_B) for c in row):
                bad.append((name, "area cell missing"))
            if not any(word in c for c in row):
                bad.append((name, "contract cell missing"))
        assert not bad, ("a supported cell of a row is gone", bad, [f["row_a"], f["row_b"]])

    def columns_aligned(t: Turn19) -> None:
        f = find(t)
        assert f, ("no table to align", t.markdown[:500])
        counts = {len(cells_of(ln)) for ln in f["table"]}
        assert len(counts) == 1, ("the lines of the table differ in their number of cells", sorted(counts), f["table"])

    def unsupported_cell_clean(t: Turn19) -> None:
        f = find(t)
        assert f, ("no table to read", t.markdown[:500])
        state = note_state(f["row_b"], t)
        leaked = [c for c in f["notes_b"] if CONDITION.search(c)]
        invented = [c for c in f["notes_a"] + f["notes_b"] if re.search(r"\d", re.sub(r"\s+", "", CITED.sub("", c)))]
        assert not leaked and not invented, ("the unsupported cell carries a figure or date the document does not hold, or the other "
                                             "group's condition", {"state_b": state, "leaked_condition": leaked, "invented": invented})

    def body(o, c, cid, rec):
        t = ask(c, cid, E2_QUESTION, rec)
        f = find(t)
        produced = bool(f and (f["notes_a"] or f["notes_b"]))
        reached = {"area_a": stated(t, AREA_A), "area_b": stated(t, AREA_B)}
        rec["info"] = base_info(t, reached, {k: t.part_stage(p) for k, p in (("area_a", AREA_A), ("area_b", AREA_B))})
        rec["info"] |= {"shape_produced": produced, "table_header": f.get("header"), "row_a": f.get("row_a"), "row_b": f.get("row_b"),
                        "note_cells_a": f.get("notes_a"), "note_cells_b": f.get("notes_b"),
                        "note_state_a": note_state(f["row_a"], t) if f else None,
                        "note_state_b": note_state(f["row_b"], t) if f else None,
                        "mechanism_exercised": bool(f) and note_state(f["row_b"], t) in ("marked_by_server", "cut_empty", "invented_let_through"),
                        "server_marks": [s["text"][:120] for s in t.spans if s["kind"] in FLAG_KINDS or SERVER_MARK.search(s["text"])],
                        "removed": [(r["action"], r["check"], r["failure"], r["text"][:160]) for r in t.removed]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "shape_produced", lambda: shape_produced(t))
        check(rec, "supported_cells_intact", lambda: supported_cells_intact(t))
        check(rec, "columns_aligned", lambda: columns_aligned(t))
        check(rec, "unsupported_cell_clean", lambda: unsupported_cell_clean(t))
        check(rec, "row_not_removed_whole", lambda: _no([t.markdown] if REMOVAL_SENTENCE.search(t.markdown) else [],
                                                        "the answer carries a removal sentence"))

    run_sequence("e2_table_row_unsupported_cell",
                 ("answered", "cited", "shape_produced", "supported_cells_intact", "columns_aligned", "unsupported_cell_clean",
                  "row_not_removed_whole"), office, body)


# --- e3: plan A's column without plan B's figures --------------------------------------------------------------------------

def test_e3_plan_as_column_holds_no_figure_of_plan_b_three_times(office, monkeypatch):
    doc = ingest_pdf(office, monkeypatch, PLANS)

    def a_stated(t: Turn19) -> None:
        missing = [n for n, pat in PLAN_FIGURES["A"].items() if not stated(t, pat)]
        assert not missing, ("a figure of plan A is not in the answer", missing, t.markdown[:500])

    def none_of_b_under_a(t: Turn19) -> None:
        wrong = [(n, ln[:200]) for n, pat in PLAN_FIGURES["B"].items() for told, ln in placements(t, pat, PLAN_WORDS) if told == "A"]
        assert not wrong, ("a figure of plan B stands under plan A", wrong[:3])

    def body(o, c, cid, rec):
        t = ask(c, cid, "מהן זכויות הבנייה לפי תכנית ג/11111 ולפי תכנית ג/22222 בנכס ברחוב האלון? הצג בטבלה עם עמודה לכל תכנית.", rec)
        reached = {f"{side}_{n}": stated(t, pat) for side, parts in PLAN_FIGURES.items() for n, pat in parts.items()}
        where = {f"{side}_{n}": [told for told, _ in placements(t, pat, PLAN_WORDS)] for side, parts in PLAN_FIGURES.items()
                 for n, pat in parts.items()}
        rec["info"] = base_info(t, reached, {f"{side}_{n}": t.part_stage(pat) for side, parts in PLAN_FIGURES.items()
                                             for n, pat in parts.items()})
        rec["info"] |= {"figure_places": where, "unknown_places": sum(p == "unknown" for v in where.values() for p in v),
                        "a_figures_under_b": [(k, p) for k, v in where.items() if k.startswith("A_") for p in v if p == "B"],
                        "removed_as_another_property": t.removed_as_another_property,
                        "tables": tables_of(t)}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "plan_a_figures_stated", lambda: a_stated(t))
        check(rec, "no_plan_b_figure_under_plan_a", lambda: none_of_b_under_a(t))

    run_sequence("e3_plan_a_column_without_plan_b_figures",
                 ("answered", "cited", "plan_a_figures_stated", "no_plan_b_figure_under_plan_a"), office, body,
                 asserted=("answered", "no_plan_b_figure_under_plan_a"))


# --- e4: a footnote condition on its own lot group only ---------------------------------------------------------------------

def test_e4_report_only_a_footnote_condition_stays_with_its_own_lot_group_three_times(office, monkeypatch):
    """REPORT ONLY: where the footnote's condition stands in the answer - with group A (its own), with group B, or with neither."""
    doc = ingest_pdf(office, monkeypatch, LOTS)

    def body(o, c, cid, rec):
        t = ask(c, cid, "מהו השטח ומהו חוזה החכירה של כל אחת מקבוצות המגרשים בשכונת הרימון, ובאיזה תנאי כפוף התשלום לפי חוזה החכירה?", rec)
        where = [(told, ln[:240]) for told, ln in placements(t, CONDITION, GROUPS)]
        sides = {told for told, _ in where}
        outcome = ("not_stated" if not where else "own_group_only" if sides == {"A"} else "on_group_b" if "B" in sides
                   else "told_to_neither")
        reached = {"area_a": stated(t, AREA_A), "area_b": stated(t, AREA_B), "condition": bool(where)}
        rec["info"] = base_info(t, reached, {"area_a": t.part_stage(AREA_A), "area_b": t.part_stage(AREA_B),
                                             "condition": t.part_stage(CONDITION)})
        rec["info"] |= {"condition_places": where, "condition_outcome": outcome,
                        "condition_on_group_b": "B" in sides, "condition_on_group_a": "A" in sides,
                        "removed_as_another_property": t.removed_as_another_property,
                        "server_span_texts": [(s["kind"], s["text"][:160]) for s in t.spans]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))

    run_sequence("e4_report_only_footnote_on_its_own_group", ("answered", "cited"), office, body, asserted=("answered",))


# --- e5: a unit titled by name and labelled by its address ------------------------------------------------------------------

def test_e5_report_only_a_unit_named_in_the_question_by_its_name_keeps_its_value_three_times(office, monkeypatch):
    """REPORT ONLY: the property named in its title block (not in its address label) is asked for by name; whether its value is kept,
    removed as another property's, or asked back (a clarification, which fails ``answered``)."""
    doc = ingest_pdf(office, monkeypatch, NAMES)
    name = NAMES["reports"][0]["name"]

    def body(o, c, cid, rec):
        t = ask(c, cid, f"מהו שווי הנכס בגישת ההשוואה של {name}?", rec)
        reached = {"value": stated(t, NAME_VALUE), "other_property_value": stated(t, OTHER_NAME_VALUE)}
        rec["info"] = base_info(t, reached, {k: t.part_stage(p) for k, p in (("value", NAME_VALUE), ("other_property_value", OTHER_NAME_VALUE))})
        rec["info"] |= {"clarification_asked": t.answer.get("status") == "clarification",
                        "removed_as_another_property": t.removed_as_another_property,
                        "removal_sentence_in_answer": bool(REMOVAL_SENTENCE.search(t.markdown))}
        check(rec, "answered", lambda: answered(t))
        check(rec, "cited", lambda: cites_document(t, doc))
        check(rec, "value_kept", lambda: _all({"value": reached["value"]}, t))
        check(rec, "no_another_property_removal", lambda: no_removal_as_another_property(t))

    run_sequence("e5_report_only_name_address_unit",
                 ("answered", "cited", "value_kept", "no_another_property_removal"), office, body, asserted=("answered",))


# --- e6: an injected wrong-plan value ---------------------------------------------------------------------------------------

@pytest.mark.skip(reason="no injection path in the real-model tests: rounds 10-19 ask the configured model and judge through the "
                         "chat API and never seed a draft (a ScriptedAgent turn replaces the model AND the judge, so a wrong-plan "
                         "value cannot be put into a real-judge turn without product code). The scripted counterpart is "
                         "tests/integration/test_round20_sequences.py (the figure of plan B under plan A is cut, the cell beside it stays).")
def test_e6_report_only_an_injected_wrong_plan_value_is_cut_or_flagged():
    raise AssertionError("skipped")


# --- the summary ----------------------------------------------------------------------------------------------------------

def build_summary() -> dict:
    turns = [t for r in RESULTS for t in r.get("turns") or []]
    purposes: dict[str, int] = {}
    statuses: dict[str, int] = {}
    for t in turns:
        for k, n in t["usage"]["purposes"].items():
            purposes[k] = purposes.get(k, 0) + n
        for k, n in t["usage"]["statuses"].items():
            statuses[k] = statuses.get(k, 0) + n
    costs = [t["usage"]["cost_usd"] for t in turns if t["usage"]["cost_usd"] is not None]
    names = list(dict.fromkeys(r["sequence"] for r in RESULTS))
    sequences = {}
    for name in names:
        s = sequence_summary(name)
        sequences[name] = {"passed": s["passed"], "passed_runs": s["passed_runs"], "shape_not_produced_runs": s["shape_not_produced_runs"],
                           "attempts": [{k: a[k] for k in ("attempt", "outcome", "checks", "failures", "stage", "info")}
                                        for a in s["attempts"]]}
    return {"model": MODEL, "turns": len(turns), "attempts": len(RESULTS), "sequences": sequences,
            "usage_purposes": purposes, "usage_statuses": statuses, "cost_usd": round(sum(costs), 4) if costs else None,
            "model_latency_s": round(sum(t["usage"]["model_latency_ms"] for t in turns) / 1000, 1),
            "wall_s": round(time.perf_counter() - STARTED, 1)}


def teardown_module(module) -> None:  # noqa: ARG001
    if RESULTS:
        flush()
        SUMMARY_FILE.write_text(json.dumps(build_summary(), ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        print(f"\nreal_model results: {os.path.relpath(RESULTS_FILE, ROOT)}\nsummary: {os.path.relpath(SUMMARY_FILE, ROOT)}")
