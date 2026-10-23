"""Round 17 acceptance with the real model (U8: R16; AE3, AE5): the three failures that U10 pins with scripted agents, shown on
the configured model and judge (``gpt-6-luna`` with the repository's settings, whose key comes from the repository's ``.env`` and is
never printed). Synthetic documents only (every name, address and figure invented), ingested through the worker into a synthetic
office of the ``rag_test`` database: the two round-17 PDFs of ``tests/fixtures/round17`` (``scripts/generate_fixtures.py --only
round17``) and one DOCX generated here.

- ``s1`` a partial list (AE3): a table of three adjustments, asked which adjustments the appraisal applied. All three reach the
  answer, each with its own citation;
- ``s2`` a scope qualifier after a table (AE5): a sensitivity table followed by a one-sentence note scoping it to one valuation
  approach. Asked for the value under a scenario, the scope (or a precise note that it was not verified) stands next to the
  value with a citation, and the value is never presented as the property's total;
- ``s3`` a scope in a section heading beside another component's qualifier (AE5): a heading names the approach of table A and a
  note under table B scopes B only. Asked for a value of table A, the scope of the heading appears and B's note is never applied.

Each sequence is asked in a **new conversation, exactly three times**; every repetition is recorded - a failure in any repetition
is reported and a later pass never hides it - and a sequence passes only when all three repetitions pass. The gating checks are by
behaviour and never by wording (a pattern part is matched, so a Hebrew definite article or prefix does not change the outcome).

Informational (recorded, not gated): the stage each part reached the answer in (``first_pass``: the model gave it; ``completed``:
the server's completion round added it from text the turn had read; ``omitted``), the completeness the turn reports, the item
ledger with the state of every item (``covered`` / ``omitted`` / ...), whether a verification (repair) round ran, the tools, the
model calls and wall time. Per turn the whole record is written to ``real_documents/eval/round17/real-model/`` (gitignored and
private: answers are kept for diagnosis), with a summary file next to it listing every repetition of every sequence.

Opt in with ``-m real_model``; skipped without a key. Host only, against ``rag_test`` (``tests/conftest.py`` guard)."""

from __future__ import annotations

import io
import json
import os
import re
import time
from datetime import UTC, datetime
from pathlib import Path

import pytest
from sqlalchemy import text

from app.config import get_settings
from app.db import tenant_tx
from tests.conftest import login
from tests.factories import make_office
from tests.integration.test_round10_real_model import new_conversation, upload, usage_of
from tests.integration.test_round10_real_model import real_key as real_key  # noqa: F401 - the key fixture

pytestmark = [pytest.mark.real_model, pytest.mark.db]

ROOT = Path(__file__).resolve().parents[3]
FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
MANIFEST = json.loads((FIXTURES / "round17" / "manifest.json").read_text(encoding="utf-8"))["documents"]
SENSITIVITY, ADJUSTMENTS = MANIFEST["sensitivity"], MANIFEST["adjustments"]
RESULTS_DIR = ROOT / "real_documents" / "eval" / "round17" / "real-model"
STAMP = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
RESULTS_FILE = RESULTS_DIR / f"results-{STAMP}.json"
SUMMARY_FILE = RESULTS_DIR / f"summary-{STAMP}.json"
RESULTS: list[dict] = []
ATTEMPTS = 3  # exactly three, each in a new conversation
STARTED = time.perf_counter()

ADMIN = "admin-r17@example.test"
MARK = "מסמך בדיקה סינתטי (סבב 17) - נתונים בדויים לבדיקת המערכת בלבד"
CITED = re.compile(r"\[([SV]\d+)\]")
MARKUP = re.compile(r"[*`‎‏‪-‮⁦-⁩]")  # markdown emphasis and bidi marks: never part of a word


# --- the synthetic documents -------------------------------------------------------------------------------------------

def docx_bytes(title: str, blocks: list[tuple[str, object]]) -> bytes:
    """A DOCX of ``blocks`` (``("h1", text)``, ``("p", text)`` or ``("table", [[cell, ...], ...])``) under ``title``, with Word
    heading styles so the reader finds the structure, and the synthetic mark as its first paragraph."""
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(MARK)
    for kind, body in blocks:
        if kind == "h1":
            doc.add_heading(str(body), 1)
        elif kind == "table":
            rows = list(body)  # type: ignore[arg-type]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, value in enumerate(cells):
                    table.cell(r, c).text = value
        else:
            doc.add_paragraph(str(body))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# s3: the approach of table A is named by its section heading alone; the note under table B scopes B only
HEADING_SCOPE = {
    "title": "שומת הדגמה למתחם אשל - שני ניתוחי רגישות", "kind": "appraisal",
    "blocks": [("h1", "1. מטרת השומה"), ("p", "שומה זו נערכה לצורך הדגמה בלבד (מסמך סינתטי)."),
               ("h1", "5. ניתוח רגישות - גישת היוון ההכנסות"),
               ("p", "הטבלה שלהלן מציגה את שווי הנכס בשינוי שיעור ההיוון."),
               ("table", [["שיעור היוון", "שווי הנכס (אלפי ₪)"], ["5.5%", "14,200"], ["6.0%", "13,600"], ["6.5%", "13,100"]]),
               ("h1", "6. ניתוח רגישות לדמי השכירות"),
               ("p", "הטבלה שלהלן מציגה את שווי הנכס בשינוי דמי השכירות."),
               ("table", [["שינוי בדמי השכירות", "שווי הנכס (אלפי ₪)"], ["-10%", "12,500"], ["0%", "13,600"], ["+10%", "14,700"]]),
               ("p", "הטבלה שלעיל מתייחסת לתרחיש של שוכר עוגן יחיד בלבד.")]}

# the scope words, with the definite article and the prefixes Hebrew puts on them (גישת ההיוון / בגישת היוון ההכנסות)
SCOPE = re.compile(r"גישת\s+ה?היוון\s+(?:של\s+)?ה?הכנסות|גישת\s+ה?הכנסות")
# a precise note that the value's scope in the source was not verified (the accepted alternative to the scope itself)
UNVERIFIED = re.compile(r"לא\s+(?:אומת|נבדק|ניתן\s+(?:לאמת|לוודא))|אינו\s+מאומת|ללא\s+אימות")
# wording that presents a figure as the property's total value
TOTAL = re.compile(r"שווי\s+(?:ה?נכס\s+)?הכולל|הערכת\s+שווי\s+כוללת|שווי\s+כולל|סך\s+(?:כל\s+)?(?:ה)?שווי")
ANCHOR_NOTE = re.compile(r"עוגן|שוכר\s+יחיד")  # the words of the note under table B (s3)


def pct(n: str) -> re.Pattern:
    """A whole percentage ``n`` followed by ``%`` (so ``5`` does not match ``15%`` or ``5.5%``), in either bidi order."""
    return re.compile(rf"(?<![\d.,]){n}(?![\d.,])\s*%")


def amount(digits: str) -> re.Pattern:
    """A figure written with or without its thousands separator (``11,800`` / ``11800``)."""
    return re.compile(r"(?<![\d.,])" + re.escape(digits[:-3]) + r"[,.]?" + re.escape(digits[-3:]) + r"(?![\d])")


# --- environment -------------------------------------------------------------------------------------------------------

@pytest.fixture
def office(db, client, monkeypatch, real_key):  # noqa: F811 - the key fixture is the imported one
    monkeypatch.setattr(get_settings(), "chat_run_inline", True)
    o = make_office(db, "משרד בדיקה סינתטי", ADMIN)
    with tenant_tx(o.ctx()) as conn:
        conn.execute(text("UPDATE office_settings SET cloud_llm_enabled = true"))
    login(client, ADMIN)
    o.client, o.ids = client, {}
    return o


def ingest_pdf(office, monkeypatch, spec: dict) -> None:
    """A round-17 PDF fixture read by the worker; ``office.ids[title]`` is its document id."""
    office.ids[spec["title"]] = upload(office, FIXTURES / spec["file"], "appraisal", monkeypatch)


def ingest_docx(office, monkeypatch, tmp_path: Path, spec: dict) -> None:
    """The synthetic DOCX ``spec`` generated and read by the worker; ``office.ids[title]`` is its document id."""
    path = tmp_path / (spec["title"].replace(" ", "_") + ".docx")
    path.write_bytes(docx_bytes(spec["title"], spec["blocks"]))
    office.ids[spec["title"]] = upload(office, path, spec["kind"], monkeypatch)


# --- turns and records -------------------------------------------------------------------------------------------------

def flush() -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    RESULTS_FILE.write_text(json.dumps({"model": "gpt-6-luna (settings)", "attempts": RESULTS}, ensure_ascii=False, indent=2,
                                       default=str), encoding="utf-8")


class Turn:
    """A finished turn: the message and answer, the checks of each verification round (its diagnostics), the item ledger the turn
    ended with and the usage."""

    def __init__(self, c, m: dict, question: str, wall: float) -> None:
        self.m, self.question = m, question
        self.answer = m.get("answer") or {}
        d = c.get(f"/api/chat/messages/{m['id']}/diagnostics")
        body = d.json() if d.status_code == 200 else {}
        self.rounds = [[{"check": p.get("check"), "kind": p.get("kind"), "component": p.get("component")} for p in r]
                       for r in body.get("rounds") or []]
        # what verification took out and why (synthetic documents only), so a part that did not reach the answer can be
        # told apart: removed by a check, or never written
        self.removed = [{"action": p.get("action"), "check": p.get("check"), "failure": p.get("failure_kind"),
                         "text": (p.get("text") or "")[:300], "reason": (p.get("reason") or "")[:300]}
                        for p in body.get("removed") or []]
        self.tools = [t.get("tool") for t in (body.get("timing") or {}).get("tools") or []]
        ledger = (body.get("resolution") or {}).get("item_ledger") or {}
        self.items = [{k: i.get(k) for k in ("id", "component", "label", "value", "state", "declared", "origin", "source",
                                             "covered_by", "why")} for i in ledger.get("items") or []]
        self.excluded = ledger.get("excluded") or []
        self.usage = usage_of(m, wall)

    @property
    def markdown(self) -> str:
        return self.answer.get("markdown") or ""

    @property
    def lines(self) -> list[str]:
        """The answer's non-empty lines without markdown emphasis and bidi marks, whitespace collapsed."""
        return [re.sub(r"\s+", " ", MARKUP.sub("", ln)).strip() for ln in self.markdown.splitlines() if ln.strip()]

    @property
    def sources(self) -> list[dict]:
        return self.answer.get("sources") or []

    @property
    def completeness(self) -> dict:
        return (self.answer.get("verification") or {}).get("completeness") or {}

    @property
    def completed_by_server(self) -> bool:
        """Whether a verification round raised a completion from read text (``REQ_TEXT_FOUND``, ``check="text_found"``)."""
        return any(p["check"] == "text_found" for r in self.rounds for p in r)

    @property
    def repair_rounds(self) -> int:
        """The verification rounds that found a problem (each one asked for a repair or an edit)."""
        return sum(1 for r in self.rounds if r)

    def documents_cited_in(self, line: str) -> set[str]:
        """The documents of the sources (S#) and of the registered values (V#) the line cites."""
        by_id = {s.get("id"): s.get("document_id") for s in [*self.sources, *(self.answer.get("values") or [])]}
        return {by_id[i] for i in CITED.findall(line) if i in by_id}

    def record(self) -> dict:
        return {"question": self.question, "message_id": self.m.get("id"), "message_status": self.m.get("status"),
                "answer_status": self.answer.get("status"), "completeness": self.completeness,
                "completed_by_server": self.completed_by_server, "repair_rounds": self.repair_rounds,
                "rounds": self.rounds, "tools": self.tools, "items": self.items, "items_excluded": self.excluded,
                "cited": [{"id": s.get("id"), "document_id": s.get("document_id"), "title": s.get("title"),
                           "location": s.get("location"), "source_kind": s.get("source_kind")} for s in self.sources],
                "usage": self.usage, "markdown": self.markdown, "removed": self.removed}


def ask(c, cid: str, question: str, rec: dict) -> Turn:
    """One turn through the API; recorded (and flushed) before anything is asserted about it."""
    started = time.perf_counter()
    r = c.post(f"/api/chat/conversations/{cid}/messages", json={"content": question})
    assert r.status_code == 200, r.text
    m = c.get(f"/api/chat/messages/{r.json()['assistant']['id']}").json()
    t = Turn(c, m, question, time.perf_counter() - started)
    rec["turns"].append(t.record())
    flush()
    assert m["status"] == "done", (m.get("status"), m.get("error"))
    return t


def check(rec: dict, name: str, fn) -> None:
    """Run one named assertion of a repetition and record its outcome; a failure never stops the repetition's other checks."""
    try:
        fn()
    except AssertionError as exc:
        rec["checks"][name] = {"outcome": "fail", "failure": str(exc)[:2000]}
    else:
        rec["checks"][name] = {"outcome": "pass"}
    flush()


def run_sequence(sequence: str, checks: tuple[str, ...], office, body) -> None:
    """``body(office, client, conversation_id, record)`` in a new conversation, exactly ``ATTEMPTS`` times; every repetition is
    recorded and none is skipped after a failure. The test fails when any check of any repetition failed (or never ran because
    a turn did not finish), with every failure in the message."""
    failures: list[str] = []
    for n in range(1, ATTEMPTS + 1):
        rec = {"sequence": sequence, "attempt": n, "turns": [], "checks": {}, "info": {}, "outcome": None}
        RESULTS.append(rec)
        try:
            body(office, office.client, new_conversation(office.client), rec)
        except AssertionError as exc:
            rec["aborted"] = str(exc)[:2000]
        for name in checks:
            rec["checks"].setdefault(name, {"outcome": "fail", "failure": "never ran: " + rec.get("aborted", "?")})
        rec["outcome"] = "pass" if all(v["outcome"] == "pass" for v in rec["checks"].values()) else "fail"
        flush()
        failures += [f"attempt {n} / {name}: {v['failure']}" for name, v in rec["checks"].items() if v["outcome"] == "fail"]
    assert not failures, "\n".join(failures)


def stage(t: Turn, reached: bool) -> str:
    """The stage the part reached the answer in: ``first_pass`` (the model gave it), ``completed`` (the server's completion round
    added it) or ``omitted``."""
    return "omitted" if not reached else "completed" if t.completed_by_server else "first_pass"


def answered(t: Turn) -> None:
    assert t.answer.get("status") != "clarification", ("a clarification was asked", t.markdown[:300])
    assert t.markdown.strip(), "an empty answer"


def cites_document(t: Turn, document: str) -> None:
    assert any(s.get("document_id") == document for s in t.sources), (
        "the document is not cited", [(s.get("id"), s.get("title")) for s in t.sources])


def has(line: str, word: str | re.Pattern) -> bool:
    return bool(word.search(line)) if isinstance(word, re.Pattern) else word in line


def near(t: Turn, anchor: str | re.Pattern, width: int = 220) -> list[str]:
    """For each place the answer states ``anchor``: its text and what follows it (to the end of the following line, else
    ``width`` characters), the stretch in which a qualifier or a citation counts as next to it."""
    lines, out = t.lines, []
    for i, ln in enumerate(lines):
        for m in (anchor.finditer(ln) if isinstance(anchor, re.Pattern) else re.finditer(re.escape(anchor), ln)):
            out.append(ln[m.start():m.start() + width] + (" " + lines[i + 1] if i + 1 < len(lines) else ""))
    return out


def around(t: Turn, anchor: str | re.Pattern) -> list[str]:
    """For each place the answer states ``anchor``: the whole line it stands in with the lines before and after it, the stretch in
    which a qualifier or a citation counts as next to a value (a scope may be written before the value or after it)."""
    lines, out = t.lines, []
    for i, ln in enumerate(lines):
        if has(ln, anchor):
            out.append(" ".join(lines[max(0, i - 1):i + 2]))
    return out


def reached_all(reached: dict[str, bool]) -> None:
    missing = [n for n, ok in reached.items() if not ok]
    assert not missing, ("the part did not reach the answer", missing)


# =========================================================================================================================
# the three sequences
# =========================================================================================================================

def test_s1_a_partial_list_reaches_the_answer_whole_with_each_citation_three_times(office, monkeypatch):
    ingest_pdf(office, monkeypatch, ADJUSTMENTS)
    doc = lambda o: o.ids[ADJUSTMENTS["title"]]  # noqa: E731
    items = {"location": ("מיקום", pct("5")), "area": ("שטח", pct("3")), "floor": ("קומה", pct("2"))}

    def cited_each(o, t: Turn) -> None:
        """Every adjustment is stated with its rate, and a citation of the adjustments' document stands next to it (its own
        line or the one that continues it)."""
        bad = []
        for name, (label, rate) in items.items():
            spans = [s for s in near(t, label, 160) if has(s, rate)]
            if not any(doc(o) in t.documents_cited_in(s) for s in spans):
                bad.append((name, spans[:2]))
        assert not bad, ("an adjustment has no citation of the appraisal next to it", bad)

    def body(o, c, cid, rec):
        t = ask(c, cid, "אילו התאמות יושמו בשומה לדירת מגורים ברחוב הדקל, ומה שיעור כל אחת?", rec)
        lines = t.lines
        reached = {name: any(has(ln, label) and has(ln, rate) for ln in lines) for name, (label, rate) in items.items()}
        rec["info"] = {"parts_reached": reached, "stage": {n: stage(t, r) for n, r in reached.items()},
                       "completeness": t.completeness.get("status"), "completed_by_server": t.completed_by_server,
                       "repair_rounds": t.repair_rounds, "tools": t.tools,
                       "ledger": [(i["label"], i["value"], i["state"], i["declared"], i["origin"]) for i in t.items]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "part", lambda: reached_all(reached))
        check(rec, "cited", lambda: cites_document(t, doc(o)))
        check(rec, "each_cited", lambda: cited_each(o, t))

    run_sequence("s1_partial_list_three_adjustments", ("answered", "part", "cited", "each_cited"), office, body)


def test_s2_a_scope_after_a_table_stands_next_to_the_value_three_times(office, monkeypatch):
    ingest_pdf(office, monkeypatch, SENSITIVITY)
    value = amount("11800")  # the 6.0% row of table 5.1

    def scoped(o, t: Turn) -> str:
        """The value is stated with the scope (or a precise unverified note) and a citation next to it, and is never put as the
        property's total; returns which form the limitation took."""
        spans = around(t, value)
        assert spans, ("the value is not stated", t.markdown[:300])
        forms = []
        for s in spans:
            form = "scope" if SCOPE.search(s) else "unverified_note" if UNVERIFIED.search(s) else None
            if form and CITED.search(s) and o.ids[SENSITIVITY["title"]] in t.documents_cited_in(s):
                forms.append(form)
        assert forms, ("neither the scope nor a precise limitation, with its citation, stands next to the value", spans[:2])
        total = [s for s in spans if TOTAL.search(s) and not SCOPE.search(s)]
        assert not total, ("the value is presented as the property's total", total[:2])
        return "scope" if "scope" in forms else forms[0]

    def body(o, c, cid, rec):
        t = ask(c, cid, "מהו שווי הנכס בשומה בשיעור היוון של 6.0%?", rec)
        lines = t.lines
        reached = {"value": any(has(ln, value) for ln in lines), "scope": any(has(ln, SCOPE) for ln in lines)}
        rec["info"] = {"parts_reached": reached, "stage": {n: stage(t, r) for n, r in reached.items()},
                       "completeness": t.completeness.get("status"), "completed_by_server": t.completed_by_server,
                       "repair_rounds": t.repair_rounds, "tools": t.tools,
                       "ledger": [(i["label"], i["value"], i["state"], i["declared"], i["origin"]) for i in t.items]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "value", lambda: reached_all({"value": reached["value"]}))
        check(rec, "cited", lambda: cites_document(t, o.ids[SENSITIVITY["title"]]))
        check(rec, "scope_next_to_value", lambda: rec["info"].update(scope_form=scoped(o, t)))

    run_sequence("s2_scope_note_after_a_table", ("answered", "value", "cited", "scope_next_to_value"), office, body)


def test_s3_a_scope_from_a_heading_appears_and_the_other_tables_note_never_does_three_times(office, monkeypatch, tmp_path):
    ingest_docx(office, monkeypatch, tmp_path, HEADING_SCOPE)
    title = HEADING_SCOPE["title"]
    value = amount("13600")  # the 6.0% row of table A (the 0% row of table B shows the same figure)

    def heading_scope(o, t: Turn) -> None:
        spans = around(t, value)
        assert spans, ("the value is not stated", t.markdown[:300])
        assert any(SCOPE.search(s) and CITED.search(s) and o.ids[title] in t.documents_cited_in(s) for s in spans), (
            "the heading's scope, with its citation, does not stand next to the value", spans[:2])

    def other_note_not_applied(t: Turn) -> None:
        hits = [ln for ln in t.lines if ANCHOR_NOTE.search(ln)]
        assert not hits, ("the note that scopes table B only is applied in the answer", hits[:2])

    def body(o, c, cid, rec):
        t = ask(c, cid, "מהו שווי הנכס בשומה למתחם אשל בשיעור היוון של 6.0%?", rec)
        lines = t.lines
        reached = {"value": any(has(ln, value) for ln in lines), "scope": any(has(ln, SCOPE) for ln in lines)}
        rec["info"] = {"parts_reached": reached, "stage": {n: stage(t, r) for n, r in reached.items()},
                       "completeness": t.completeness.get("status"), "completed_by_server": t.completed_by_server,
                       "repair_rounds": t.repair_rounds, "tools": t.tools,
                       "ledger": [(i["label"], i["value"], i["state"], i["declared"], i["origin"]) for i in t.items]}
        check(rec, "answered", lambda: answered(t))
        check(rec, "value", lambda: reached_all({"value": reached["value"]}))
        check(rec, "cited", lambda: cites_document(t, o.ids[title]))
        check(rec, "scope_from_heading", lambda: heading_scope(o, t))
        check(rec, "other_note_not_applied", lambda: other_note_not_applied(t))

    run_sequence("s3_scope_in_heading_beside_other_note",
                 ("answered", "value", "cited", "scope_from_heading", "other_note_not_applied"), office, body)


# --- the summary -------------------------------------------------------------------------------------------------------

def build_summary() -> dict:
    turns = [t for r in RESULTS for t in r.get("turns") or []]
    purposes: dict[str, int] = {}
    statuses: dict[str, int] = {}
    for t in turns:
        for k, n in t["usage"]["purposes"].items():
            purposes[k] = purposes.get(k, 0) + n
        for k, n in t["usage"]["statuses"].items():
            statuses[k] = statuses.get(k, 0) + n
    costs = [t["usage"]["cost_usd"] for t in turns if t["usage"]["cost_usd"] is not None]
    sequences: dict[str, list[dict]] = {}
    for r in RESULTS:
        info = r.get("info") or {}
        sequences.setdefault(r["sequence"], []).append(
            {"attempt": r["attempt"], "outcome": r["outcome"], "checks": {k: v["outcome"] for k, v in r["checks"].items()},
             "failures": {k: v["failure"] for k, v in r["checks"].items() if v["outcome"] == "fail"},
             "parts_reached": info.get("parts_reached"), "stage": info.get("stage"), "completeness": info.get("completeness"),
             "repair_rounds": info.get("repair_rounds"), "scope_form": info.get("scope_form"), "ledger": info.get("ledger")})
    return {"model": "gpt-6-luna (settings)", "turns": len(turns), "attempts": len(RESULTS),
            "sequences": {k: {"attempts": v, "passed": len(v) == ATTEMPTS and all(a["outcome"] == "pass" for a in v),
                              "passed_runs": f"{sum(a['outcome'] == 'pass' for a in v)}/{len(v)}"}
                          for k, v in sequences.items()},
            "usage_purposes": purposes, "usage_statuses": statuses, "cost_usd": round(sum(costs), 4) if costs else None,
            "model_latency_s": round(sum(t["usage"]["model_latency_ms"] for t in turns) / 1000, 1),
            "wall_s": round(time.perf_counter() - STARTED, 1)}


def teardown_module(module) -> None:  # noqa: ARG001
    if RESULTS:
        flush()
        SUMMARY_FILE.write_text(json.dumps(build_summary(), ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nreal_model results: {os.path.relpath(RESULTS_FILE, ROOT)}\nsummary: {os.path.relpath(SUMMARY_FILE, ROOT)}")
