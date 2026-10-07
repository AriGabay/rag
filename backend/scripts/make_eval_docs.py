"""Writes the synthetic documents of evaluation set v5 to tests/fixtures/eval_v5/: four small appraisal reports
with invented places, addresses and values (no real office data).

They give the evaluation situations the real reports alone do not: three reports in one invented town (a
question about the town is about a set), a survey table with seven asking rents (an answer about the table must
keep its multiplicity), a sentence with a value per m² and a monthly rent per m² (the two must stay apart), a
value on an equivalent-area basis including VAT, a negative statement, and a land report in another town.

The file names are the documents' titles (a title is taken from the uploaded file's name, as in the office).
Run from backend/: ``uv run python scripts/make_eval_docs.py``.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "eval_v5"
TOWN = "קריית אלון"


def report(name: str, title: str, sections: list[tuple[str, list]]) -> None:
    doc = Document()
    doc.add_heading(title, level=0)
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        for item in body:
            if isinstance(item, str):
                doc.add_paragraph(item)
            else:  # a table: (caption, header, rows, notes)
                caption, header, rows, notes = item
                doc.add_paragraph(caption)
                t = doc.add_table(rows=1, cols=len(header))
                for c, h in zip(t.rows[0].cells, header, strict=True):
                    c.text = h
                for r in rows:
                    cells = t.add_row().cells
                    for c, v in zip(cells, r, strict=True):
                        c.text = v
                for n in notes:
                    doc.add_paragraph(n)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT / name)


def main() -> None:
    report("שומה - הדקל 14 קריית אלון.docx", f"שומת מקרקעין – בניין משרדים, רחוב הדקל 14, {TOWN}", [
        ("1. מטרת חוות הדעת", [f"חוות דעת זו נערכה לצורך הערכת שווי השוק של בניין משרדים ברחוב הדקל 14, {TOWN}, "
                               "לצורכי בטוחה."]),
        ("2. תיאור הנכס", ["הבניין בן 6 קומות משרדים מעל 2 קומות חניה תת-קרקעיות, ובו 48 מקומות חניה.",
                           "השטח הבנוי ברוטו למשרדים הוא 4,200 מ\"ר."]),
        ("3. סקר שוק", [("להלן נתוני היצע לשכירות משרדים בסביבת הנכס:",
                         ["כתובת", "קומה", "שטח במ\"ר", "דמ\"ש למ\"ר לחודש"],
                         [["הדקל 2", "3", "180", "₪ 55"], ["הדקל 9", "5", "240", "₪ 58"],
                          ["האורן 4", "2", "150", "₪ 60"], ["האורן 11", "7", "320", "₪ 63"],
                          ["הברוש 6", "4", "210", "₪ 61"], ["הברוש 15", "1", "130", "₪ 57"],
                          ["שדרות הנחל 20", "6", "400", "₪ 68"]],
                         ["(*) המחירים המבוקשים אינם כוללים מע\"מ ודמי ניהול."])]),
        ("4. השומה", ["בהתאם לסקר ולמיקום הנכס, ראוי לקבוע את השווי למ\"ר בנוי ברוטו למשרדים בגבולות של 11,500 ₪, "
                      "ללא מע\"מ, ודמי שכירות ראויים בגבולות של 62 ₪ למ\"ר לחודש.",
                      "השווי הכולל של הנכס הוא 48,300,000 ₪, ללא מע\"מ."]),
    ])
    report("שומה - שדרות הנחל 8 קריית אלון.docx", f"שומת מקרקעין – בניין מגורים, שדרות הנחל 8, {TOWN}", [
        ("1. מטרת חוות הדעת", ["הערכת שווי דירות בבניין מגורים חדש לצורכי ליווי בנקאי."]),
        ("2. תיאור הנכס", ["בבניין 36 יחידות דיור, מהן 12 דירות נשואות חוות הדעת.",
                           "בבניין 40 מקומות חניה ומחסן לכל דירה."]),
        ("3. השומה", ["השווי למ\"ר אקוו' לדירות בבניין נקבע בכ-24,000 ₪ למ\"ר, כולל מע\"מ.",
                      "שווי 12 הדירות נשואות חוות הדעת הוא 18,600,000 ₪, כולל מע\"מ."]),
        ("4. מצב משפטי", ["לא רשומות הערות אזהרה על הנכס, ולא נמצאו שעבודים."]),
    ])
    report("שומה - התמר 3 קריית אלון.docx", f"שומת מקרקעין – מבנה מסחרי, רחוב התמר 3, {TOWN}", [
        ("1. תיאור הנכס", ["מבנה מסחרי בן קומה אחת, בשטח בנוי של 860 מ\"ר, ובו 9 חנויות."]),
        ("2. השומה", ["השווי למ\"ר בנוי למסחר נקבע ל-9,800 ₪, ללא מע\"מ.",
                      "שיעור התפוסה במבנה הוא 89%."]),
        ("3. היטל השבחה", ["לא צפויה חבות בהיטל השבחה בגין הנכס."]),
    ])
    report("חוות דעת - מתחם הגורן נווה צור.docx", "חוות דעת – מתחם הגורן, נווה צור", [
        ("1. תיאור המקרקעין", ["מתחם קרקע בשטח 14 דונם, בייעוד תעשייה ומלאכה."]),
        ("2. השומה", ["שווי דונם קרקע במתחם נקבע ל-1,200,000 ₪.",
                      "ערך הכינון של המבנים הקיימים במתחם הוא כ-7.5 מיליון ₪."]),
    ])


if __name__ == "__main__":
    main()
