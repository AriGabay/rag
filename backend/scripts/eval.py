"""Run the evaluation set against the LIVE demo stack and write docs/evaluation/eval-results.md (U15).

    cd backend && uv run python scripts/eval.py
    # options: --questions eval/questions.yaml --out ../docs/evaluation/eval-results.md --only N01,T04 --json out.json

The held-out general set (U12; eval/questions_general.yaml, scored by eval/general.py) runs with
``--set general``; ``--real-sample`` additionally requires office A in cloud mode (it refuses otherwise and
never changes settings) and writes docs/evaluation/real-model-sample.md. Items whose extraction was still
running in background jobs are asked again once the job queue drains; both runs are recorded.
``--rescore results.json`` scores stored answers again without asking anything.

    cd backend && uv run python scripts/eval.py --set general --real-sample --json /tmp/general.json

Needs the Compose stack seeded with scripts/seed_demo.py (BACKEND_URL, default http://localhost:8000;
demo password from DEMO_PASSWORD, default demo1234). Expected values are computed from
tests/fixtures/ground_truth.yaml (eval/truth.py) for each item's conditions, the asking user's groups,
and the live review queue: records the queue still lists are not verified and are excluded.

Table-extraction and field metrics read the live database read-only through the runtime role in each
office's system context (EVAL_DATABASE_URL, default built from the repository .env). They are
reported separately from answer quality.

The report states which providers answered: with an office in demo mode content answers come from the
deterministic mock (not a cloud-model benchmark); with an office in cloud mode they come from the real model.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import platform
import re
import subprocess
import sys
import tempfile
import time
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

import httpx
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))  # before scripts/, so `eval` is the package, not this file

from eval.flows import ask_flow  # noqa: E402
from eval.truth import (  # noqa: E402
    EXTRACTION_SQL_OCCURRENCES,
    EXTRACTION_SQL_TABLES,
    MIXED_KEYS,
    USERS,
    Filters,
    by_filename,
    docs,
    expected_clarifications,
    expected_stats,
    field_accuracy,
    gt_record_for,
    price_token,
    table_accuracy,
    title_of,
    unverified_from_queue,
)

BACKEND = os.environ.get("BACKEND_URL", "http://localhost:8000")
PASSWORD = os.environ.get("DEMO_PASSWORD", "demo1234")
REPO = ROOT.parent
GROUP_NAMES = {"שומות רמת גן וגבעתיים": "G1", "שומות פרויקטים": "G2", "ידע כללי": "G3"}  # scripts/seed_demo.py
ANSWERED = ("numeric", "combined", "content")
NUMERIC_FIELDS = ("record_count", "mean_price_per_sqm", "weighted_price_per_sqm", "median_price_per_sqm",
                  "min_price_per_sqm", "max_price_per_sqm")


def pct(values: list[float], q: float) -> float | None:
    if not values:
        return None
    s = sorted(values)
    k = (len(s) - 1) * q
    lo, hi = int(k), min(int(k) + 1, len(s) - 1)
    return s[lo] + (s[hi] - s[lo]) * (k - lo)


def rate(ok: int, n: int) -> str:
    return f"{ok}/{n} ({100 * ok / n:.1f}%)" if n else "n/a"


def big_numbers(s: str) -> list[str]:
    out = []
    for m in re.findall(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?", s or ""):
        v = float(m.replace(",", ""))
        if v >= 1000 and not (1950 <= v <= 2100 and "," not in m and "." not in m):
            out.append(m)
    return out


class Live:
    """Logged-in clients and the live document / verification state."""

    def __init__(self):
        self.clients: dict[str, httpx.Client] = {}
        self.versions: dict[str, dict[str, str]] = {}  # document id -> {version id: filename}
        self.files: dict[str, bytes] = {}

    def client(self, email: str) -> httpx.Client:
        if email not in self.clients:
            c = httpx.Client(base_url=BACKEND, timeout=120)
            r = c.post("/api/auth/login", json={"email": email, "password": PASSWORD})
            r.raise_for_status()
            self.clients[email] = c
        return self.clients[email]

    def load(self) -> None:
        a, b = self.client("admin-a@demo.test"), self.client("admin-b@demo.test")
        self.docs_a = [d for d in a.get("/api/documents").json()["documents"] if not d["deleted"]]
        self.docs_b = [d for d in b.get("/api/documents").json()["documents"] if not d["deleted"]]
        self.ids_a = {d["id"] for d in self.docs_a}
        self.ids_b = {d["id"] for d in self.docs_b}
        self.group_of = {d["id"]: GROUP_NAMES.get(d["group"]["name"], d["group"]["name"]) for d in self.docs_a}
        # titles per office: browser tests may upload copies of office A fixtures into office B
        self.by_title = {"A": {d["title"]: d["id"] for d in self.docs_a},
                         "B": {d["title"]: d["id"] for d in self.docs_b}}
        unverified_a, unmapped_a = unverified_from_queue(a, "A")
        unverified_b, unmapped_b = unverified_from_queue(b, "B")
        self.unverified = unverified_a | unverified_b
        self.unmapped = unmapped_a + unmapped_b
        self.office_ids = {k: self.client(e).get("/api/auth/me").json()["office"]["id"]
                           for k, e in (("A", "admin-a@demo.test"), ("B", "admin-b@demo.test"))}

    def doc_id(self, gt_doc: str) -> str | None:
        return self.by_title[docs()[gt_doc]["office"]].get(title_of(gt_doc))

    def forbidden(self, user: str, rule: str | None) -> set[str]:
        office, groups = USERS[user]
        other_office = self.ids_b if office == "A" else self.ids_a
        if rule == "other_office":
            return other_office
        if rule == "other_group":
            return other_office | {d for d in self.ids_a if groups and self.group_of.get(d) not in groups}
        return set()

    def filename(self, client: httpx.Client, document_id: str, version_id: str) -> str | None:
        if document_id not in self.versions:
            r = self.client("admin-a@demo.test").get(f"/api/documents/{document_id}")
            if r.status_code != 200:
                r = self.client("admin-b@demo.test").get(f"/api/documents/{document_id}")
            self.versions[document_id] = {v["id"]: v["filename"] for v in r.json().get("versions", [])}
        return self.versions[document_id].get(version_id)

    def file(self, client: httpx.Client, url: str) -> tuple[int, str, bytes]:
        base = url.partition("#")[0]
        if base not in self.files:
            r = client.get(base)
            self.files[base] = (r.status_code, r.headers.get("content-type", ""), r.content)
        return self.files[base]


def page_text(content_type: str, data: bytes, page: int | None) -> str:
    if content_type.startswith("application/pdf"):
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if page is None or page > len(pdf.pages):
                return ""
            return pdf.pages[page - 1].extract_text() or ""
    import docx  # DOCX: the whole document (no physical pages)

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def evidence_check(live: Live, client, answer: dict, data_kind: str) -> tuple[int, int, list[str]]:
    """Record sources (not retrieved passages) whose file page shows the record's price."""
    ok = n = 0
    problems = []
    for s in answer.get("sources", []):
        if s.get("chunk_id"):
            continue  # a retrieved passage of a combined answer, not a calculation source
        n += 1
        filename = live.filename(client, s["document_id"], s["version_id"])
        row = s.get("row")
        rec = gt_record_for(filename or "", 0 if row is not None else None, row, data_kind)
        status, ctype, data = live.file(client, s["url"])
        if rec is None or status != 200:
            problems.append(f"source {s.get('title')} p{s.get('page_list')} row {row}: "
                            f"{'unmapped' if rec is None else f'HTTP {status}'}")
            continue
        page = (s.get("page_list") or [None])[0]
        if docs()[rec["_doc"]]["kind"] == "pdf_scanned":
            ok += 1  # image-only page: the stored OCR text is checked by the acceptance suite
            continue
        token = price_token(rec["id"])
        text_ = re.sub(r"\s+", " ", page_text(ctype, data, page))
        if token in text_ and (page == rec["page"] or rec["page"] is None):
            ok += 1
        else:
            problems.append(f"{rec['id']}: price {token} not on cited page {page} (answer key page {rec['page']})")
    return ok, n, problems


def run_item(live: Live, item: dict) -> list[dict]:
    user = item["user"]
    client = live.client(user)
    office, groups = USERS[user]
    results = []
    conversation = None
    previous: Filters | None = None
    for i, turn in enumerate(item["turns"], start=1):
        expect = turn["expect"]
        filters = Filters.of(expect["filters"]) if expect.get("filters") else None
        if "clarify" in turn:
            answers = turn["clarify"] or {}
        elif filters:
            answers = {k: getattr(filters, k) for k in ("data_kind", "date_field", *MIXED_KEYS)
                       if getattr(filters, k) is not None}
        else:
            answers = {}
        started = time.perf_counter()
        flow = ask_flow(client, turn["ask"], answers, conversation)
        total_ms = (time.perf_counter() - started) * 1000
        conversation = flow.conversation_id
        a = flow.answer
        kind = a.get("kind", "error")
        reasons: list[str] = []
        res = {"id": item["id"], "turn": i, "category": item["category"], "user": user, "question": turn["ask"],
               "expected": expect["outcome"], "got": kind, "provider": a.get("provider"),
               "cached": bool(a.get("cached")), "latencies_ms": flow.latencies_ms, "total_ms": total_ms,
               "requests_cached": [bool(x.get("cached")) for x in flow.transcript],
               "clarifications": flow.clarifications, "sources": len(a.get("sources", []))}
        if flow.error:
            reasons.append(flow.error)
        out = expect["outcome"]
        numeric = a.get("numeric")
        if out in ("numeric", "combined"):
            exp = expected_stats(filters, office, groups, live.unverified)
            res["expected_numeric"] = exp.as_answer()
            if exp.count == 0:
                reasons.append("SPEC: no verified answer-key records match these filters")
            if kind != out:
                reasons.append(f"kind {kind} (expected {out})")
            elif numeric:
                got = {k: numeric.get(k) for k in NUMERIC_FIELDS}
                res["got_numeric"] = got
                diff = {k: (got[k], exp.as_answer()[k]) for k in NUMERIC_FIELDS if got[k] != exp.as_answer()[k]}
                if diff:
                    reasons.append("numbers differ (got, expected): " + json.dumps(diff, ensure_ascii=False))
                res["numeric_exact"] = not diff
            else:
                reasons.append("no numeric block")
            # a follow-up inherits the conditions confirmed in the conversation
            inherited = tuple(k for k in MIXED_KEYS if previous is not None and getattr(previous, k) is not None)
            want_clar = expected_clarifications(filters, tuple(expect.get("stated", [])) + inherited, office,
                                                groups, live.unverified)
            asked = [k for k in flow.clarifications if k in MIXED_KEYS]
            res["clarification_precise"] = asked == want_clar
            if asked != want_clar:
                res["clarification_note"] = f"asked {asked}, needed {want_clar}"
        if kind in ("numeric", "combined") and numeric:
            ok, n, problems = evidence_check(live, client, a, (filters.data_kind if filters else
                                                               "transaction_price"))
            res["evidence"] = (ok, n)
            res["evidence_problems"] = problems
        if out == "clarification":
            if kind != "clarification":
                reasons.append(f"kind {kind} (expected a clarification about {expect['key']})")
            elif a["clarification"]["key"] != expect["key"]:
                reasons.append(f"asked about {a['clarification']['key']} (expected {expect['key']})")
        if out == "abstain" and kind != "abstain":
            reasons.append(f"kind {kind} (expected abstain)")
        if out in ("clarification", "abstain") and (numeric or big_numbers(a.get("text", ""))):
            reasons.append("a number was shown")
            res["number_shown"] = True
        if out == "content" and kind != "content":
            reasons.append(f"kind {kind} (expected content)")
        if expect.get("sources"):
            want = [(live.doc_id(s["doc"]), s["page"]) for s in expect["sources"]]
            hit = any(s["document_id"] == d and (p is None or p in s.get("page_list", []))
                      for s in a.get("sources", []) for d, p in want)
            res["source_recall"] = hit
            if not hit:
                reasons.append("expected source not cited: " + ", ".join(
                    f"{s['doc']} p.{s['page']}" for s in expect["sources"]) + " | got: " + ", ".join(
                    f"{s['title'][:3].strip()} p{s.get('page_list')}" for s in a.get("sources", [])))
        if expect.get("forbid"):
            bad = {s["document_id"] for s in a.get("sources", [])} & live.forbidden(user, expect["forbid"])
            res["isolation_ok"] = not bad
            if bad:
                reasons.append(f"LEAK: {len(bad)} source(s) outside the user's scope")
        res["ok"] = not reasons
        res["reasons"] = reasons
        res["text"] = (a.get("text") or a.get("detail") or "")[:240]
        results.append(res)
        if filters is not None:
            previous = filters
    return results


def runtime_db_url() -> str | None:
    """The live database through the runtime role (RLS applies): EVAL_DATABASE_URL, else built from .env."""
    url = os.environ.get("EVAL_DATABASE_URL")
    if url:
        return url
    env = {}
    env_file = REPO / ".env"
    if env_file.exists():
        for line in env_file.read_text().splitlines():
            if "=" in line and not line.lstrip().startswith("#"):
                k, v = line.split("=", 1)
                env[k.strip()] = v.strip()
    if not env.get("RAG_APP_PASSWORD"):
        return None
    return f"postgresql+psycopg://rag_app:{env['RAG_APP_PASSWORD']}@localhost:5433/rag"


def extraction_metrics(live: Live) -> dict:
    url = runtime_db_url()
    if not url:
        return {"error": "no EVAL_DATABASE_URL and no RAG_APP_PASSWORD in .env"}
    from sqlalchemy import create_engine, text

    tables, occurrences, models = [], [], Counter()
    try:
        engine = create_engine(url)
        for office, office_id in live.office_ids.items():
            mine = {d["filename"] for d in by_filename().values() if d["office"] == office}
            with engine.begin() as conn:
                conn.execute(text("SELECT set_config('app.office_id', :o, true), set_config('app.user_id', '', true),"
                                  " set_config('app.role', 'system', true)"), {"o": office_id})
                tables += [dict(r._mapping) for r in conn.execute(text(EXTRACTION_SQL_TABLES)) if r.filename in mine]
                occurrences += [dict(r._mapping) for r in conn.execute(text(EXTRACTION_SQL_OCCURRENCES))
                                if r.filename in mine]
                for m, n in conn.execute(text("SELECT embedding_model, count(*) FROM chunks GROUP BY 1")):
                    models[m] += n
        engine.dispose()
    except Exception as exc:  # noqa: BLE001
        return {"error": f"database not reachable: {type(exc).__name__}"}
    return {"tables": table_accuracy(tables), "fields": field_accuracy(occurrences), "embedding_models": dict(models)}


def environment(live: Live) -> dict:
    def sh(cmd: list[str]) -> str:
        try:
            return subprocess.run(cmd, capture_output=True, text=True, timeout=20).stdout.strip()
        except Exception:  # noqa: BLE001
            return ""

    env = {"date": datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC"), "backend": BACKEND,
           "python": platform.python_version(), "os": f"{platform.system()} {platform.release()} {platform.machine()}"}
    if sys.platform == "darwin":
        env["cpu"] = sh(["sysctl", "-n", "machdep.cpu.brand_string"])
        env["cores"] = sh(["sysctl", "-n", "hw.ncpu"])
        mem = sh(["sysctl", "-n", "hw.memsize"])
        env["memory"] = f"{int(mem) / 2**30:.0f} GiB" if mem.isdigit() else mem
    docker = sh(["docker", "info", "--format", "{{.ServerVersion}}|{{.NCPU}}|{{.MemTotal}}|{{.OperatingSystem}}"])
    if docker:
        v, ncpu, memt, os_ = (docker.split("|") + ["", "", "", ""])[:4]
        env["docker"] = f"Docker {v} on {os_}, {ncpu} CPUs, {int(memt) / 2**30:.1f} GiB" if memt.isdigit() else docker
    env["commit"] = sh(["git", "-C", str(REPO), "rev-parse", "--short", "HEAD"])
    env["health"] = live.client("admin-a@demo.test").get("/api/health").json()
    for key, email in (("A", "admin-a@demo.test"), ("B", "admin-b@demo.test")):
        s = live.client(email).get("/api/admin/settings").json()
        env[f"provider_{key}"] = s.get("mode")
        env[f"model_{key}"] = s.get("model")
    env["demo_mode"] = live.client("admin-a@demo.test").get("/api/auth/me").json().get("demo_mode")
    return env


def summarize(results: list[dict]) -> dict:
    m: dict = {}
    m["turns"] = len(results)
    m["items"] = len({r["id"] for r in results})
    m["correct"] = sum(r["ok"] for r in results)
    num = [r for r in results if r["expected"] in ("numeric", "combined")]
    m["numeric"] = (sum(bool(r.get("numeric_exact")) for r in num), len(num))
    neg = [r for r in results if r["expected"] in ("abstain", "clarification")]
    m["justified"] = (sum(r["got"] in ("abstain", "clarification") and not r.get("number_shown") for r in neg), len(neg))
    m["exact_negative"] = (sum(r["ok"] for r in neg), len(neg))
    answerable = [r for r in results if r["expected"] in ANSWERED]
    m["answered"] = (sum(r["got"] in ANSWERED for r in answerable), len(answerable))
    m["unjustified_abstain"] = (sum(r["got"] in ("abstain", "clarification") for r in answerable), len(answerable))
    rec = [r for r in results if "source_recall" in r]
    m["source_recall"] = (sum(r["source_recall"] for r in rec), len(rec))
    ev = [r for r in results if "evidence" in r]
    m["evidence_answers"] = (sum(1 for r in ev if r["evidence"][0] == r["evidence"][1] and r["evidence"][1] > 0), len(ev))
    m["evidence_sources"] = (sum(r["evidence"][0] for r in ev), sum(r["evidence"][1] for r in ev))
    iso = [r for r in results if "isolation_ok" in r]
    m["isolation"] = (sum(r["isolation_ok"] for r in iso), len(iso))
    cp = [r for r in results if "clarification_precise" in r]
    m["clarification_precision"] = (sum(r["clarification_precise"] for r in cp), len(cp))
    lat = [ms for r in results for ms in r["latencies_ms"]]
    m["latency_requests"] = lat
    m["latency_turns"] = [r["total_ms"] for r in results]
    pairs = [(ms, c) for r in results for ms, c in zip(r["latencies_ms"], r["requests_cached"], strict=False)]
    m["latency_fresh"] = [ms for ms, c in pairs if not c]
    m["latency_cached"] = [ms for ms, c in pairs if c]
    m["providers"] = Counter(r["provider"] for r in results)
    m["cached"] = sum(r["cached"] for r in results)
    by_cat = defaultdict(lambda: [0, 0])
    for r in results:
        by_cat[r["category"]][0] += r["ok"]
        by_cat[r["category"]][1] += 1
    m["by_category"] = dict(sorted(by_cat.items()))
    return m


def provider_note(env: dict, m: dict) -> str:
    """Which providers answered, stated honestly: a mock/extractive run is not a cloud-model benchmark."""
    modes = f"office A mode `{env.get('provider_A')}`, office B mode `{env.get('provider_B')}`"
    seen = dict(m["providers"])
    if "cloud" in seen:
        model = env.get("model_A") or "n/a"
        return (f"> **Mixed providers.** {modes}; model `{model}`; providers seen in answers: {seen}. Office A "
                "questions the rules do not fully explain are interpreted and composed by the real cloud model; "
                "fully explained price questions stay SQL + templates with no model call. Office B answers come "
                "from the labeled demo mock. Latencies include real network/model time for cloud turns. All data "
                "is synthetic (R36); quality on real appraisal reports is **pending** until authorized real "
                "documents are supplied.")
    return (f"> **Not a cloud-model benchmark.** The demo stack answers content questions with the deterministic "
            f"mock provider ({modes}; providers seen in answers: {seen}). Numeric answers are SQL + templates with "
            "no model call. Latencies below are local mock latencies, not cloud-provider latencies. All data is "
            "synthetic (R36); quality on real appraisal reports is **pending** until authorized real documents are "
            "supplied.")


def fmt_ms(v: float | None) -> str:
    return "n/a" if v is None else f"{v:.0f} ms"


def write_report(path: Path, env: dict, m: dict, results: list[dict], extraction: dict, live: Live,
                 elapsed: float) -> None:
    lat, turns = m["latency_requests"], m["latency_turns"]
    lines = [
        "# Evaluation results (U15)",
        "",
        f"Generated by `cd backend && uv run python scripts/eval.py` on {env['date']} against the live demo "
        f"stack ({env['backend']}), commit `{env.get('commit') or 'n/a'}`. Run time {elapsed:.0f} s.",
        "",
        provider_note(env, m),
        "",
        "## Environment",
        "",
        "| item | value |",
        "|---|---|",
        f"| machine | {env.get('cpu', 'n/a')}, {env.get('cores', '?')} cores, {env.get('memory', '?')} RAM ({env['os']}) |",
        f"| Docker | {env.get('docker', 'n/a')} |",
        f"| embedding model in the stack (chunks.embedding_model) | {', '.join(f'`{k}` ({v} chunks)' for k, v in (extraction.get('embedding_models') or {}).items()) or 'n/a'} |",
        f"| answer provider (mode) | office A `{env.get('provider_A')}` (`{env.get('model_A')}`), office B "
        f"`{env.get('provider_B')}` (demo mode: {env.get('demo_mode')}) |",
        f"| database role / RLS bypass | `{env['health'].get('db_role')}` / {env['health'].get('db_bypass_rls')} |",
        f"| evaluation set | {m['items']} items, {m['turns']} scored turns (`backend/eval/questions.yaml`) |",
        f"| verification state | {len(live.unverified)} answer-key records still in the review queue and excluded "
        f"from expectations: {', '.join(sorted(live.unverified)) or 'none'} |",
        "",
        "## Answer quality",
        "",
        "| metric | result | definition |",
        "|---|---|---|",
        f"| answer correctness (all turns) | {rate(m['correct'], m['turns'])} | outcome, numbers, sources and isolation all as expected |",
        f"| numeric exact match | {rate(*m['numeric'])} | count, mean, weighted, median, min, max equal the ground-truth oracle (Decimal, ROUND_HALF_UP) |",
        f"| justified abstention / clarification | {rate(*m['justified'])} | questions without basis that got an abstention or clarification and no number |",
        f"| ... with the exact expected behaviour | {rate(*m['exact_negative'])} | right kind (and the right clarification key) |",
        f"| answered rate (coverage) | {rate(*m['answered'])} | answerable questions that got a numeric/content answer (a system refusing everything scores 0) |",
        f"| unjustified abstention / clarification | {rate(*m['unjustified_abstain'])} | answerable questions left at an abstention or an open clarification |",
        f"| source recall (content) | {rate(*m['source_recall'])} | expected document + physical page among the cited sources |",
        f"| evidence coverage (numeric answers) | {rate(*m['evidence_answers'])} | answers whose every record source opens and shows the record's price on the cited page |",
        f"| evidence coverage (sources) | {rate(*m['evidence_sources'])} | individual record sources that pass that check |",
        f"| isolation | {rate(*m['isolation'])} | turns with a scope rule and no source from another office / group |",
        f"| clarification precision | {rate(*m['clarification_precision'])} | numeric turns where the system asked exactly the mixed-attribute clarifications the verified data requires |",
        f"| latency per request p50 / p95 / max | {fmt_ms(pct(lat, .5))} / {fmt_ms(pct(lat, .95))} / {fmt_ms(max(lat) if lat else None)} | {len(lat)} `/api/ask` calls, incl. clarification answers |",
        f"| ... computed (not from cache) p50 / p95 | {fmt_ms(pct(m['latency_fresh'], .5))} / {fmt_ms(pct(m['latency_fresh'], .95))} | {len(m['latency_fresh'])} requests, incl. every clarification step (never cached) |",
        f"| ... served from cache p50 / p95 | {fmt_ms(pct(m['latency_cached'], .5))} / {fmt_ms(pct(m['latency_cached'], .95))} | {len(m['latency_cached'])} requests |",
        f"| latency per turn p50 / p95 | {fmt_ms(pct(turns, .5))} / {fmt_ms(pct(turns, .95))} | question plus its clarification round-trips |",
        f"| answers served from cache | {m['cached']}/{m['turns']} | exact-condition cache hits (R28); repeated runs hit more |",
        "",
        "### By category",
        "",
        "| category | correct |",
        "|---|---|",
    ]
    lines += [f"| {cat} | {rate(ok, n)} |" for cat, (ok, n) in m["by_category"].items()]
    lines += ["", "## Table and field extraction (separate from answer quality)", ""]
    if "error" in extraction:
        lines.append(f"Not measured: {extraction['error']}.")
    else:
        t, f = extraction["tables"], extraction["fields"]
        lines += [
            "Compared with `ground_truth.yaml` for every current document version in both offices (read-only, "
            "runtime role, system context).",
            "",
            "| metric | result |",
            "|---|---|",
            f"| table header accuracy | {rate(*t['headers'])} |",
            f"| table cell accuracy | {rate(*t['cells'])} |",
            f"| row physical-page accuracy | {rate(*t['row_pages'])} |",
            f"| critical-field accuracy (data kind, city, price, area, area type, transaction/valuation date) | {rate(*f['critical'])} |",
            f"| all-field accuracy (17 fields per record) | {rate(*f['all'])} |",
            f"| records extracted / expected | {f['records_found']}/{f['records_expected']} (extra occurrences: {f['extra']}) |",
            "",
            "| document | headers | cells | row pages |",
            "|---|---|---|---|",
        ]
        lines += [f"| {d} | {v['headers']} | {v['cells']} | {v['row_pages']} |" for d, v in t["per_doc"].items()]
        if f["missing"] or f["critical_errors"]:
            lines += ["", "Field errors:", ""]
            lines += [f"- missing record: {x}" for x in f["missing"]]
            lines += [f"- {x}" for x in f["critical_errors"][:40]]
    failures = [r for r in results if not r["ok"]]
    lines += ["", f"## Failures ({len(failures)})", ""]
    if not failures:
        lines.append("None.")
    for r in failures:
        lines.append(f"- **{r['id']}** turn {r['turn']} ({r['category']}, {r['user']}): «{r['question']}» — "
                     f"expected `{r['expected']}`, got `{r['got']}`. " + "; ".join(r["reasons"]))
    notes = [r for r in results if r.get("clarification_note") or r.get("evidence_problems")]
    if notes:
        lines += ["", "## Notes (not counted as failures)", ""]
        for r in notes:
            if r.get("clarification_note"):
                lines.append(f"- {r['id']} turn {r['turn']}: clarifications {r['clarification_note']}")
            for p in r.get("evidence_problems", []):
                lines.append(f"- {r['id']} turn {r['turn']}: evidence — {p}")
    lines += ["", "## All turns", "", "| id | turn | category | expected | got | provider | ok | ms |", "|---|---|---|---|---|---|---|---|"]
    lines += [f"| {r['id']} | {r['turn']} | {r['category']} | {r['expected']} | {r['got']} | {r['provider']} | "
              f"{'yes' if r['ok'] else '**no**'} | {r['total_ms']:.0f} |" for r in results]
    lines += [
        "",
        "## How to read this",
        "",
        "- Expected numbers are recomputed from the answer key for each item's conditions; records still in the "
        "review queue are excluded, so a correct system must not count them.",
        "- The cache is keyed on conditions and data version; questions asked earlier (by this script, the browser "
        "tests or a person) can be served from it, which shortens latencies and can mask a cache defect.",
        "- Real-document quality evaluation is pending: run the same script against a stack seeded with authorized "
        "real reports and a manually verified answer key in the same YAML format.",
        "",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


# --- the held-out general set (U12, KTD16) ------------------------------------------------------------

class LiveReader:
    """Read-only access to stored turns, provider usage and the job queue through the runtime role, in each
    office's system context (rag_app cannot read another user's questions otherwise)."""

    def __init__(self, live: Live):
        from sqlalchemy import create_engine

        url = runtime_db_url()
        if not url:
            raise SystemExit("--set general needs EVAL_DATABASE_URL or RAG_APP_PASSWORD in .env to read stored plans")
        self.engine = create_engine(url)
        self.office_ids = live.office_ids

    def _rows(self, office: str, sql: str, **params) -> list[dict]:
        from sqlalchemy import text

        with self.engine.begin() as conn:
            conn.execute(text("SELECT set_config('app.office_id', :o, true), set_config('app.user_id', '', true),"
                              " set_config('app.role', 'system', true)"), {"o": self.office_ids[office]})
            return [dict(r._mapping) for r in conn.execute(text(sql), params)]

    def plan_of(self, question_id: str) -> dict | None:
        for office in ("A", "B"):
            rows = self._rows(office, "SELECT plan, steps, parse_route FROM questions WHERE id = CAST(:q AS uuid)",
                              q=question_id)
            if rows:
                return rows[0]
        return None

    def usage_since(self, office: str, since: datetime) -> list[dict]:
        return self._rows(office, "SELECT provider, model, purpose, status, ok, count(*) AS calls,"
                          " coalesce(sum(input_tokens), 0) AS input_tokens,"
                          " coalesce(sum(output_tokens), 0) AS output_tokens, round(avg(latency_ms)) AS avg_ms"
                          " FROM provider_usage WHERE created_at >= :t GROUP BY 1, 2, 3, 4, 5 ORDER BY 3, 4",
                          t=since)

    def open_extract_jobs(self, office: str = "A") -> int:
        return self._rows(office, "SELECT count(*) AS n FROM jobs WHERE kind = 'extract_facts'"
                          " AND status IN ('queued', 'running')")[0]["n"]


def general_hooks(live: Live, reader: LiveReader, cloud_off: bool = False):
    from eval.general import Hooks, filename_to_doc

    admin = live.client("admin-a@demo.test")

    def doc_of(source: dict) -> str | None:
        name = live.filename(admin, source["document_id"], source["version_id"])
        return filename_to_doc().get(name or "")

    return Hooks(client=live.client, doc_of=doc_of, plan_of=reader.plan_of, forbidden=live.forbidden,
                 unverified=live.unverified, cloud_off=cloud_off)


def wait_for_extraction(reader: LiveReader, timeout: float = 600) -> bool:
    deadline = time.time() + timeout
    while time.time() < deadline:
        if reader.open_extract_jobs() == 0:
            return True
        time.sleep(5)
    return False


def print_general_line(r: dict) -> None:
    if r.get("not_run"):
        print(f"skip {r['id']}    {r['not_run']}", flush=True)
        return
    mark = "ok " if r["ok"] else "FAIL"
    print(f"{mark} {r['id']}.{r['turn']:<2} {r['expected']:>13} -> {r['got'] or '-':<13} task {r['task_type']}"
          f" {r['total_ms']:6.0f} ms" + ("" if r["ok"] else "  " + "; ".join(r["reasons"])[:300]), flush=True)


def run_general(args) -> int:
    from eval.general import load_items, rescore, run_and_score, summarize

    live = Live()
    live.load()
    settings_a = live.client("admin-a@demo.test").get("/api/admin/settings").json()
    if args.real_sample and not args.rescore and settings_a.get("mode") != "cloud":
        print(f"REFUSED: --real-sample needs office A in cloud mode; it is in `{settings_a.get('mode')}` "
              f"(status {settings_a.get('mode_status')}). Enable cloud use for office A in the admin screen and run "
              "the connection test first; this script never changes office settings.", file=sys.stderr)
        return 2
    reader = LiveReader(live)
    items = load_items(args.questions_general)
    if args.only:
        wanted = set(args.only.split(","))
        items = [i for i in items if i["id"] in wanted]
    hooks = general_hooks(live, reader)
    if args.rescore:  # score stored answers again with the current scorer; nothing is asked
        data = json.loads(Path(args.rescore).read_text(encoding="utf-8"))
        env, elapsed, usage = data["env"], data["elapsed_s"], data["usage"]
        settings_a = data.get("settings_a") or settings_a
        wanted_ids = {i["id"] for i in items}
        final = rescore(hooks, items, [r for r in data["results"] if r["id"] in wanted_ids])
        first_runs = {k: rescore(hooks, items, v) for k, v in data.get("first_runs", {}).items() if k in wanted_ids}
        for r in final:
            print_general_line(r)
    else:
        env = environment(live)
        since = datetime.now(UTC)
        started = time.perf_counter()
        results = run_and_score(hooks, items, on_turn=print_general_line)
        # an item whose extraction was still running in background jobs is asked again once the worker drains
        pending = sorted({r["id"] for r in results if not r.get("not_run") and (r.get("pending_extraction") or 0) > 0})
        reruns: dict[str, list[dict]] = {}
        if pending:
            print(f"\n{len(pending)} items reported pending extraction: {', '.join(pending)}; waiting for the worker")
            drained = wait_for_extraction(reader)
            print("worker drained" if drained else "worker still busy after the timeout; asking again anyway")
            again = run_and_score(hooks, [i for i in items if i["id"] in pending], on_turn=print_general_line)
            for r in again:
                reruns.setdefault(r["id"], []).append(r)
        order = [i["id"] for i in items]
        final = [r for r in results if r["id"] not in reruns] + [r for v in reruns.values() for r in v]
        final.sort(key=lambda r: (order.index(r["id"]), r["turn"]))
        first_runs = {k: [r for r in results if r["id"] == k] for k in reruns}
        elapsed = time.perf_counter() - started
        usage = {o: reader.usage_since(o, since) for o in ("A", "B")}
    m = summarize(final)
    print_general_summary(m)
    out_json = Path(args.json or (Path(tempfile.gettempdir()) / "rag-eval-general.json"))
    out_json.write_text(json.dumps({"env": env, "settings_a": {k: settings_a.get(k) for k in (
        "provider", "provider_name", "model", "mode", "mode_status", "untested")}, "summary": m,
        "results": final, "first_runs": first_runs, "usage": usage, "elapsed_s": elapsed},
        ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(f"results: {out_json}")
    if args.real_sample:
        out = Path(args.out_sample)
        write_real_sample(out, env, settings_a, m, final, first_runs, usage, elapsed)
        print(f"report: {out}")
        gate = real_sample_gate(m)
        print("REAL-MODEL GATE: " + ("PASS" if gate["pass"] else "FAIL") + f" ({gate['reason']})")
        return 0 if gate["pass"] else 1
    return 0


def print_general_summary(m: dict) -> None:
    ok, n = m["held_out"]
    print(f"\nitems passed {m['passed']}/{m['run']} (not run: {', '.join(m['not_run']) or 'none'}) | held-out "
          f"{rate(ok, n)} | turns {rate(*m['turns'])} | task type {rate(*m['task_type'])} | cache hits {m['cached']}")
    print("\ncategory               passed")
    for cat, (p, total) in m["by_category"].items():
        print(f"{cat:<22} {p}/{total}")
    for tag, v in m["ae"].items():
        print(f"{tag}: passed {v['passed']} of {v['items']}" + (f", not run {v['not_run']}" if v["not_run"] else ""))


def real_sample_gate(m: dict) -> dict:
    """R27: every AE1-AE6 item that ran passes and the held-out pass rate is at least 80%."""
    ok, n = m["held_out"]
    ae_fail = sorted(i for tag, v in m["ae"].items() if tag in {f"AE{k}" for k in range(1, 7)}
                     for i in v["items"] if i not in v["passed"] and i not in v["not_run"])
    held = ok / n if n else 0.0
    reasons = []
    if ae_fail:
        reasons.append(f"AE items failed: {', '.join(ae_fail)}")
    if held < 0.8:
        reasons.append(f"held-out pass rate {100 * held:.1f}% < 80%")
    return {"pass": not reasons, "reason": "; ".join(reasons) or f"held-out {100 * held:.1f}%, all AE items that ran "
            "passed", "ae_failed": ae_fail, "held_out_rate": held}


def _cell(s) -> str:
    return str(s if s is not None else "—").replace("|", "\\|").replace("\n", " ")


def write_real_sample(path: Path, env: dict, settings: dict, m: dict, results: list[dict], first_runs: dict,
                      usage: dict, elapsed: float) -> None:
    gate = real_sample_gate(m)
    ok, n = m["held_out"]
    lines = [
        "# Real-model sample (U12, R27)",
        "",
        f"Generated by `cd backend && uv run python scripts/eval.py --set general --real-sample` on {env['date']} "
        f"against the live stack ({env['backend']}), commit `{env.get('commit') or 'n/a'}`. Run time {elapsed:.0f} s.",
        "",
        "| item | value |",
        "|---|---|",
        f"| provider / model | {settings.get('provider_name')} / `{settings.get('model')}` |",
        f"| office A mode | `{settings.get('mode')}` (connection test pending: {settings.get('untested')}) |",
        "| documents sent to the model | synthetic fixtures only (`backend/tests/fixtures/`, `general/`); no real "
        "appraisal content, no key or secret in this report |",
        "| question set | `backend/eval/questions_general.yaml` (written before any prompt tuning; not edited after "
        "seeing results) |",
        f"| items | {m['items']} ({m['run']} run; not run: {', '.join(m['not_run']) or 'none'}) |",
        "",
        "## Gate",
        "",
        f"**{'PASS' if gate['pass'] else 'FAIL'}** — {gate['reason']}.",
        "",
        f"- Held-out items passed: {rate(ok, n)} (threshold 80%, an assumption of the plan).",
        f"- All items passed: {m['passed']}/{m['run']}; scored turns {rate(*m['turns'])}; plan task type as expected "
        f"{rate(*m['task_type'])}; answers served from cache: {m['cached']}.",
        "",
        "| acceptance example | items | passed | not run |",
        "|---|---|---|---|",
    ]
    lines += [f"| {tag} | {', '.join(v['items'])} | {', '.join(v['passed']) or '—'} | {', '.join(v['not_run']) or '—'} |"
              for tag, v in m["ae"].items()]
    lines += ["", "## Failures by scored facet", "",
              "Items failing each facet of the question-file rules (an item can fail several). Items that fail "
              f"only on citation precision: {', '.join(m['only_precision']) or 'none'}.", "",
              "| facet | items failing |", "|---|---|"]
    lines += [f"| {facet} | {n} |" for facet, n in m["facet_failures"].items()]
    lines += ["", "## By category", "", "| category | passed |", "|---|---|"]
    lines += [f"| {cat} | {p}/{t} |" for cat, (p, t) in m["by_category"].items()]
    lines += ["", "## Per item", "",
              "| id | turn | category | held-out | expected | got | task type (expected) | tools | pass |",
              "|---|---|---|---|---|---|---|---|---|"]
    for r in results:
        if r.get("not_run"):
            lines.append(f"| {r['id']} | 1 | {r['category']} | {'yes' if r['held_out'] else 'no'} | {r['expected']} |"
                         f" — | — | — | not run |")
            continue
        lines.append(f"| {r['id']} | {r['turn']} | {r['category']} | {'yes' if r['held_out'] else 'no'} | "
                     f"{r['expected']} | {r['got']}{' (cache)' if r['cached'] else ''} | {r['task_type']} "
                     f"({r.get('expected_task_type') or '—'}) | {', '.join(t for t in r['tools'] if t) or '—'} | "
                     f"{'yes' if r['ok'] else '**no**'} |")
    failures = [r for r in results if r.get("ok") is False]
    lines += ["", f"## Failures ({len({r['id'] for r in failures})} items)", ""]
    for r in failures:
        lines.append(f"- **{r['id']}** turn {r['turn']} ({r['category']}): «{r['question']}» — expected "
                     f"`{r['expected']}`, got `{r['got']}`. " + "; ".join(_cell(x) for x in r["reasons"]))
    if first_runs:
        lines += ["", "## Re-asked after background extraction", "",
                  "These items reported documents still pending extraction; they were asked again once the worker "
                  "drained. The table above holds the second run; the first run was:", ""]
        for k, rs in first_runs.items():
            lines += [f"- {k} turn {r['turn']}: got `{r['got']}`, pending extraction {r.get('pending_extraction')}, "
                      f"{'pass' if r['ok'] else 'fail: ' + '; '.join(r['reasons'])[:300]}" for r in rs]
    lines += ["", "## Provider calls during the run", "",
              "| office | provider | model | purpose | status | calls | input tokens | output tokens | avg ms |",
              "|---|---|---|---|---|---|---|---|---|"]
    for office, rows in usage.items():
        lines += [f"| {office} | {u['provider']} | {u['model']} | {u['purpose']} | {u['status'] or ('ok' if u['ok'] else '?')}"
                  f" | {u['calls']} | {u['input_tokens']} | {u['output_tokens']} | {u['avg_ms']} |" for u in rows]
    lines += [
        "",
        "## How this is scored",
        "",
        "Scoring follows the header of `backend/eval/questions_general.yaml` (implemented in "
        "`backend/eval/general.py`); wording is never scored. The plan facets (task type, relation, tool) come "
        "from the stored turn. A computed figure counts the reviewed figure, else the separately labeled "
        "preliminary figure (model-extracted values stay preliminary until a person reviews them). Items that "
        "need cloud use off (AE2) cannot run against an office in cloud mode; they are proven by the scripted "
        "gate 8 and the U9 tests instead.",
        "",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--questions", default=str(ROOT / "eval" / "questions.yaml"))
    ap.add_argument("--out", default=str(REPO / "docs" / "evaluation" / "eval-results.md"))
    ap.add_argument("--only", default="")
    ap.add_argument("--json", default="")
    ap.add_argument("--set", default="existing", choices=("existing", "general"),
                    help="existing: eval/questions.yaml (77 turns); general: eval/questions_general.yaml (held-out)")
    ap.add_argument("--questions-general", default=str(ROOT / "eval" / "questions_general.yaml"))
    ap.add_argument("--real-sample", action="store_true",
                    help="with --set general: require office A in cloud mode and write the real-model sample report")
    ap.add_argument("--out-sample", default=str(REPO / "docs" / "evaluation" / "real-model-sample.md"))
    ap.add_argument("--rescore", default="", help="with --set general: score a stored results JSON again (no questions)")
    args = ap.parse_args(argv)
    if args.real_sample and args.set != "general":
        ap.error("--real-sample runs the general set: add --set general")
    if args.set == "general":
        return run_general(args)

    items = yaml.safe_load(Path(args.questions).read_text(encoding="utf-8"))["items"]
    if args.only:
        wanted = set(args.only.split(","))
        items = [i for i in items if i["id"] in wanted]
    started = time.perf_counter()
    live = Live()
    live.load()
    env = environment(live)
    results = []
    for item in items:
        for r in run_item(live, item):
            results.append(r)
            mark = "ok " if r["ok"] else "FAIL"
            print(f"{mark} {r['id']}.{r['turn']:<2} {r['expected']:>13} -> {r['got']:<13} {r['total_ms']:6.0f} ms"
                  + ("" if r["ok"] else "  " + "; ".join(r["reasons"])[:200]), flush=True)
    m = summarize(results)
    extraction = extraction_metrics(live)
    elapsed = time.perf_counter() - started
    write_report(Path(args.out), env, m, results, extraction, live, elapsed)
    if args.json:
        Path(args.json).write_text(json.dumps({"env": env, "results": results}, ensure_ascii=False, indent=1,
                                              default=str), encoding="utf-8")
    lat = m["latency_requests"]
    print(f"\ncorrect {rate(m['correct'], m['turns'])} | numeric exact {rate(*m['numeric'])} | justified "
          f"abstention {rate(*m['justified'])} | answered {rate(*m['answered'])} | source recall "
          f"{rate(*m['source_recall'])} | evidence {rate(*m['evidence_answers'])} | isolation {rate(*m['isolation'])}"
          f" | p50 {fmt_ms(pct(lat, .5))} p95 {fmt_ms(pct(lat, .95))}")
    if "error" not in extraction:
        print(f"tables: cells {rate(*extraction['tables']['cells'])}, headers {rate(*extraction['tables']['headers'])}"
              f" | critical fields {rate(*extraction['fields']['critical'])}")
    print(f"report: {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
